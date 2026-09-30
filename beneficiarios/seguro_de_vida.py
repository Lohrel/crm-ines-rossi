"""Planilha do seguro de vida, no layout que o escritório envia à seguradora (Tokio Marine).

Traz todas as vidas do contrato no dia 1º do mês, com o salário, o capital
(o salário vezes o múltiplo do contrato, até o capital máximo) e a fatura de
cada vida (o capital vezes a taxa por mil). O escritório usa o total da fatura
para conferir a cobrança da seguradora.

As colunas ficam nas mesmas posições da planilha do escritório, inclusive as
vazias, para o arquivo seguir direto para a seguradora. Capital e fatura saem
como fórmulas: quem preencher no Excel um salário que faltava vê tudo recalcular.
"""

import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from functools import cached_property
from io import BytesIO

from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apolices.models import Contrato, Subfatura
from beneficiarios.models import Vinculo

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CENTAVOS = Decimal("0.01")
# O escritório recebe os dados perto do dia 30 e envia a planilha à seguradora até o dia 5.
DIA_LIMITE_DE_ENVIO = 5
A_PARTIR_DO_DIA = 20
REAIS = '_-"R$" * #,##0.00_-;-"R$" * #,##0.00_-;_-"R$" * "-"??_-;_-@_-'
FONTE = Font(name="Tahoma", size=8)
NEGRITO = Font(name="Tahoma", size=8, bold=True)


@dataclass(frozen=True)
class Coluna:
    letra: str
    titulo: str
    largura: float
    formato: str = "General"


# Posições da planilha do escritório (as colunas sem título ficam vazias e estreitas).
LOCAL = Coluna("A", "Local", 9.5)
SUB = Coluna("B", "IDENTIFICADOR DE SUB", 20)
EMPRESA = Coluna("C", "Empresa", 31)
NOME = Coluna("E", "Nome", 35)
NASCIMENTO = Coluna("H", "Nascimento", 14.5, "DD/MM/YYYY")
CPF = Coluna("N", "CPF", 20)
SALARIO = Coluna("O", "Valor Salário", 12, "#,##0.00")
EMAIL = Coluna("P", "E-MAIL", 12)
DDD = Coluna("Q", "DDD", 9)
FONE = Coluna("R", "FONE", 13)
MULTIPLO = Coluna("S", "{}x", 14, REAIS)
CAPITAL = Coluna("T", "Capital Segurado", 14.5, REAIS)
FATURA = Coluna("U", "Fatura", 12, REAIS)
COLUNAS = [
    LOCAL,
    SUB,
    EMPRESA,
    NOME,
    NASCIMENTO,
    CPF,
    SALARIO,
    EMAIL,
    DDD,
    FONE,
    MULTIPLO,
    CAPITAL,
    FATURA,
]
ULTIMA = FATURA.letra
LARGURA_DAS_VAZIAS = 3.3


def mes_sugerido(hoje: date) -> date:
    """Do dia 20 em diante, a planilha do mês seguinte; antes, a do mês corrente."""
    if hoje.day < A_PARTIR_DO_DIA:
        return hoje.replace(day=1)
    return date(hoje.year + hoje.month // 12, hoje.month % 12 + 1, 1)


def telefone_em_partes(telefone: str) -> tuple[str, str]:
    """DDD e número, como nas colunas DDD e FONE (sem DDD, o número vai inteiro)."""
    digitos = re.sub(r"\D", "", telefone or "")
    if digitos.startswith("0"):
        digitos = digitos[1:]
    if digitos.startswith("55") and len(digitos) > 11:
        digitos = digitos[2:]
    if len(digitos) >= 10:
        return digitos[:2], digitos[2:]
    return "", digitos


def ordem_do_codigo(codigo: str) -> tuple:
    """Códigos numéricos em ordem de número (2 antes de 10); os outros, em ordem alfabética."""
    return (0, int(codigo), "") if codigo.isdigit() else (1, 0, codigo)


@dataclass
class Vida:
    vinculo: Vinculo
    subfatura: Subfatura | None

    @property
    def pessoa(self):
        return self.vinculo.beneficiario

    @property
    def local(self) -> str:
        return self.subfatura.local if self.subfatura else ""

    @property
    def identificador(self) -> str:
        return self.subfatura.identificador if self.subfatura else ""

    @property
    def boleto_separado(self) -> bool:
        return bool(self.subfatura and self.subfatura.boleto_separado)

    @property
    def empresa(self) -> str:
        unidade = self.pessoa.unidade_efetiva
        return unidade.razao_social if unidade else self.pessoa.cliente.nome

    @property
    def salario(self) -> Decimal | None:
        return self.vinculo.salario

    @property
    def capital(self) -> Decimal | None:
        return self.vinculo.capital

    @property
    def fatura(self) -> Decimal | None:
        return self.vinculo.fatura


@dataclass
class Subtotal:
    """Fatura de parte das vidas, com as linhas delas na planilha (para a fórmula da soma)."""

    rotulo: str
    vidas: list[Vida]
    linhas: list[int]

    @property
    def fatura(self) -> Decimal:
        return _soma_das_faturas(self.vidas).quantize(CENTAVOS)

    def formula(self) -> str:
        faixas = []
        for linha in self.linhas:
            if faixas and faixas[-1][1] == linha - 1:
                faixas[-1][1] = linha
            else:
                faixas.append([linha, linha])
        partes = [
            f"{FATURA.letra}{inicio}" + (f":{FATURA.letra}{fim}" if fim > inicio else "")
            for inicio, fim in faixas
        ]
        return f"=SUM({','.join(partes)})"


def _soma_das_faturas(vidas: list[Vida]) -> Decimal:
    """Soma sem arredondar cada vida: como no Excel, arredonda-se só o valor mostrado."""
    return sum((vida.fatura or 0 for vida in vidas), Decimal(0))


class PlanilhaDoSeguroDeVida:
    """Vidas de um contrato de seguro de vida no dia 1º do mês, com capital e fatura."""

    def __init__(self, contrato: Contrato, mes: date) -> None:
        self.contrato = contrato
        self.data = mes.replace(day=1)
        self.vidas = self._vidas()

    def _vidas(self) -> list[Vida]:
        vinculos = (
            Vinculo.objects.ativos(self.data)
            .filter(contrato=self.contrato)
            .select_related(
                "beneficiario__cliente",
                "beneficiario__unidade",
                "beneficiario__titular__unidade",
            )
        )
        subfaturas = self.contrato.subfaturas_por_unidade()
        vidas = []
        for vinculo in vinculos:
            # Capital e fatura saem sempre das regras deste contrato.
            vinculo.contrato = self.contrato
            unidade = vinculo.beneficiario.unidade_efetiva
            vidas.append(Vida(vinculo, subfaturas.get(unidade.pk) if unidade else None))
        vidas.sort(key=lambda v: (ordem_do_codigo(v.identificador), v.empresa, v.pessoa.nome))
        return vidas

    @property
    def prazo_de_envio(self) -> date:
        return self.data.replace(day=DIA_LIMITE_DE_ENVIO)

    @property
    def titulo_da_aba(self) -> str:
        return f"seguro de vida {self.data:%d.%m.%Y}"

    @property
    def nome_do_arquivo(self) -> str:
        cliente = slugify(self.contrato.cliente.nome)
        return f"seguro-de-vida-{cliente}-{self.data:%Y-%m}.xlsx"

    # Resumo para conferir a cobrança na tela, sem abrir o Excel.

    @cached_property
    def sem_salario(self) -> list[Vida]:
        if not self.contrato.usa_salario:
            return []
        return [vida for vida in self.vidas if vida.salario is None]

    @cached_property
    def sem_subfatura(self) -> list[Vida]:
        """Vidas de filiais sem subfatura cadastrada neste contrato (Local e SUB em branco)."""
        if not self.contrato.cliente.eh_empresa:
            return []
        return [vida for vida in self.vidas if not vida.identificador]

    @property
    def capital_total(self) -> Decimal:
        return sum((vida.capital or 0 for vida in self.vidas), Decimal(0))

    @property
    def fatura_total(self) -> Decimal | None:
        if self.contrato.taxa_por_mil is None:
            return None
        return _soma_das_faturas(self.vidas).quantize(CENTAVOS)

    @cached_property
    def boletos_separados(self) -> list[Subtotal]:
        """Uma soma por empresa das filiais que pagam boleto à parte."""
        if self.contrato.taxa_por_mil is None:
            return []
        grupos: dict[str, Subtotal] = {}
        for linha, vida in enumerate(self.vidas, start=2):
            if vida.boleto_separado:
                rotulo = f"Boleto separado: {vida.empresa}"
                grupo = grupos.setdefault(vida.empresa, Subtotal(rotulo, [], []))
                grupo.vidas.append(vida)
                grupo.linhas.append(linha)
        return [grupos[empresa] for empresa in sorted(grupos)]

    @property
    def fatura_sem_os_separados(self) -> Decimal | None:
        if not self.boletos_separados:
            return None
        separadas = [vida for grupo in self.boletos_separados for vida in grupo.vidas]
        restante = _soma_das_faturas(self.vidas) - _soma_das_faturas(separadas)
        return restante.quantize(CENTAVOS)

    # Planilha

    def gerar(self) -> bytes:
        livro = Workbook()
        folha = livro.active
        folha.title = self.titulo_da_aba
        self._cabecalho(folha)
        for numero, vida in enumerate(self.vidas, start=2):
            self._linha(folha, numero, vida)
        self._totais(folha, ultima_vida=len(self.vidas) + 1)
        folha.freeze_panes = "A2"
        folha.auto_filter.ref = f"A1:{ULTIMA}{max(len(self.vidas) + 1, 2)}"
        arquivo = BytesIO()
        livro.save(arquivo)
        return arquivo.getvalue()

    def _cabecalho(self, folha) -> None:
        for numero in range(1, 22):
            folha.column_dimensions[get_column_letter(numero)].width = LARGURA_DAS_VAZIAS
        for coluna in COLUNAS:
            titulo = coluna.titulo.format(self.contrato.capital_multiplo or "")
            celula = folha[f"{coluna.letra}1"]
            celula.value = titulo if coluna is not MULTIPLO or self.contrato.usa_salario else None
            celula.font = NEGRITO
            folha.column_dimensions[coluna.letra].width = coluna.largura

    def _linha(self, folha, numero: int, vida: Vida) -> None:
        pessoa = vida.pessoa
        ddd, fone = telefone_em_partes(pessoa.telefone or pessoa.whatsapp)
        valores = {
            LOCAL: _numero_se_puder(vida.local),
            SUB: _numero_se_puder(vida.identificador),
            EMPRESA: vida.empresa,
            NOME: pessoa.nome,
            NASCIMENTO: pessoa.data_nascimento,
            CPF: pessoa.cpf_formatado or None,
            SALARIO: vida.salario,
            EMAIL: pessoa.email or None,
            DDD: _numero_se_puder(ddd),
            FONE: _numero_se_puder(fone),
            MULTIPLO: None,
            CAPITAL: None,
            FATURA: None,
        }
        contrato = self.contrato
        if contrato.usa_salario:
            valores[MULTIPLO] = f"={SALARIO.letra}{numero}*{contrato.capital_multiplo}"
            valores[CAPITAL] = f"={MULTIPLO.letra}{numero}"
            if contrato.capital_maximo is not None:
                maximo = _no_excel(contrato.capital_maximo)
                valores[CAPITAL] = f"=MIN({MULTIPLO.letra}{numero},{maximo})"
        else:
            valores[CAPITAL] = contrato.capital_maximo
        if contrato.taxa_por_mil is not None:
            taxa = _no_excel(contrato.taxa_por_mil)
            valores[FATURA] = f"=({CAPITAL.letra}{numero}*{taxa})/1000"
        for coluna, valor in valores.items():
            celula = folha[f"{coluna.letra}{numero}"]
            celula.value = valor
            celula.font = FONTE
            celula.number_format = coluna.formato
            if isinstance(valor, str) and valor.startswith("=") and coluna not in _FORMULAS:
                # Texto que começa com "=" (ex.: num nome) fica como texto, nunca como fórmula.
                celula.data_type = "s"

    def _totais(self, folha, ultima_vida: int) -> None:
        """Total da fatura; com boleto separado, o valor dele e o do restante."""
        if self.contrato.taxa_por_mil is None or not self.vidas:
            return
        total = ultima_vida + 1
        self._total(folha, total, "Total", f"=SUM({FATURA.letra}2:{FATURA.letra}{ultima_vida})")
        separados = self.boletos_separados
        for linha, grupo in enumerate(separados, start=total + 1):
            self._total(folha, linha, grupo.rotulo, grupo.formula())
        if separados:
            linha = total + len(separados) + 1
            menos = "".join(f"-{FATURA.letra}{total + i}" for i in range(1, len(separados) + 1))
            formula = f"={FATURA.letra}{total}{menos}"
            self._total(folha, linha, "Total sem o boleto separado", formula)

    @staticmethod
    def _total(folha, linha: int, rotulo: str, formula: str) -> None:
        folha[f"{CAPITAL.letra}{linha}"] = rotulo
        folha[f"{CAPITAL.letra}{linha}"].font = NEGRITO
        folha[f"{CAPITAL.letra}{linha}"].data_type = "s"
        folha[f"{FATURA.letra}{linha}"] = formula
        folha[f"{FATURA.letra}{linha}"].font = NEGRITO
        folha[f"{FATURA.letra}{linha}"].number_format = REAIS


_FORMULAS = {MULTIPLO, CAPITAL, FATURA}


def _numero_se_puder(codigo: str) -> int | str | None:
    # Na planilha do escritório, Local e SUB são números; com zero à esquerda, fica o texto.
    if not codigo:
        return None
    return int(codigo) if codigo.isdigit() and not codigo.startswith("0") else codigo


def _no_excel(valor: Decimal) -> str:
    """Número para a fórmula: ponto decimal, sem zeros sobrando e sem notação 1E+2."""
    return format(valor.normalize(), "f")
