"""Importação da planilha do seguro de vida, no layout que o escritório envia à seguradora.

Carrega de uma vez as vidas de um contrato de seguro de vida: acha cada pessoa
pelo CPF ou pelo nome (ou cadastra o titular novo), coloca no plano com o salário
e guarda a subfatura (Local e SUB) de cada filial. Não cria movimentações: é o
retrato do que já existe, como o "Lançar plano atual" da página da pessoa.

Como na importação da planilha mensal, a prévia faz tudo numa transação desfeita
no fim, para mostrar exatamente o que vai acontecer; ao confirmar, faz de novo,
agora gravando. Importar a mesma planilha outra vez não repete nada: só o que
mudou (ex.: um salário) é atualizado.
"""

import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from functools import cached_property

from django.core import signing
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import DatabaseError, transaction
from django.db.models import Q
from django.utils.formats import number_format
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apolices.models import Contrato, Subfatura
from beneficiarios.models import Beneficiario, Vinculo
from beneficiarios.seguro_de_vida import CENTAVOS, ordem_do_codigo
from clientes.documentos import CPF
from clientes.models import Cliente, Unidade
from clientes.planilhas import (
    NAO_LEU,
    PlanilhaInvalidaError,
    chave,
    conferir_tamanho,
    mesmo_codigo,
)

# Nome de cada coluna da planilha (sem acento, em maiúsculas) e o dado que ela traz.
COLUNAS = {
    "LOCAL": "local",
    "IDENTIFICADOR DE SUB": "sub",
    "SUB": "sub",
    "SUBFATURA": "sub",
    "EMPRESA": "empresa",
    "NOME": "nome",
    "NASCIMENTO": "nascimento",
    "DATA DE NASCIMENTO": "nascimento",
    "DATA NASCIMENTO": "nascimento",
    "CPF": "cpf",
    "VALOR SALARIO": "salario",
    "SALARIO": "salario",
    "E MAIL": "email",
    "EMAIL": "email",
    "DDD": "ddd",
    "FONE": "fone",
    "TELEFONE": "fone",
    "CELULAR": "fone",
    "CAPITAL SEGURADO": "capital",
    "CAPITAL": "capital",
    "FATURA": "fatura",
}
DADOS = set(COLUNAS.values())
EM_REAIS = {"salario", "capital", "fatura"}
LIMITE_DE_LINHAS = 2000
# O .xlsx é um .zip: descompactado, a planilha de vidas fica bem abaixo disto.
LIMITE_DESCOMPACTADO = 20 * 1024 * 1024
VALIDADE_DA_PREVIA = timedelta(hours=2)
ASSINATURA = "beneficiarios.importacao_de_vidas"
# Até um centavo de diferença entre a planilha e a conta do contrato é arredondamento.
TOLERANCIA = Decimal("0.01")

# O que acontece com cada vida.
IMPORTAR = "importar"
ATUALIZAR = "atualizar"
PULAR = "pular"
FILIAL = "filial"
PROBLEMA = "problema"


# ----------------------------------------------------------------------
# Leitura do arquivo
# ----------------------------------------------------------------------


@dataclass
class Linha:
    numero: int
    valores: dict[str, str]

    def __getitem__(self, nome: str) -> str:
        # Só os dados das colunas; o resto (ex.: numero, no template) continua atributo.
        if nome not in DADOS:
            raise KeyError(nome)
        return self.valores.get(nome, "")


@dataclass
class PlanilhaLida:
    arquivo: str
    aba: str
    linhas: list[Linha]
    avisos: list[str]

    @property
    def data(self) -> date | None:
        """A data do nome da aba (ex.: "seguro de vida 01.05.2026"), se tiver."""
        return data_no_titulo(self.aba)


def data_no_titulo(titulo: str) -> date | None:
    achado = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", titulo)
    if not achado:
        return None
    dia, mes, ano = (int(parte) for parte in achado.groups())
    try:
        return date(ano, mes, dia)
    except ValueError:
        return None


def texto_do_valor(valor) -> str:
    """Valor em reais com ponto decimal (ex.: 2500.5); o que não for número fica como veio."""
    if isinstance(valor, int | float) and not isinstance(valor, bool):
        return str(Decimal(str(valor)))
    texto = re.sub(r"(?i)r\$|\s", "", str(valor))
    if not texto.strip("-"):
        return ""  # vazio, ou o "-" do zero no formato contábil
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?|\d+,\d+", texto):
        return texto.replace(".", "").replace(",", ".")  # como se escreve no Brasil: 2.500,50
    return texto


def _texto(valor, nome: str) -> str:
    if valor is None:
        return ""
    if nome == "nascimento":
        if isinstance(valor, int | float) and 3000 < valor < 80000:
            valor = from_excel(valor)  # data guardada como número
        if isinstance(valor, datetime | date):
            return f"{valor:%d/%m/%Y}"
    if nome in EM_REAIS:
        return texto_do_valor(valor)
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if nome == "cpf" and isinstance(valor, int):
        return f"{valor:011d}"  # CPF guardado como número perde o zero da frente
    return " ".join(str(valor).split())


def _cabecalho(linhas: list[tuple]) -> tuple[int, dict[int, str]] | None:
    """A linha dos títulos (Nome, CPF, Valor Salário...) entre as primeiras da aba."""
    for numero, valores in enumerate(linhas, start=1):
        colunas = {i: COLUNAS[chave(v)] for i, v in enumerate(valores) if chave(v) in COLUNAS}
        if "nome" in colunas.values() and len(colunas) >= 3:
            return numero, colunas
    return None


def ler_planilha(arquivo) -> PlanilhaLida:
    """As vidas da planilha, com avisos sobre o que ficou de fora."""
    conferir_tamanho(arquivo, LIMITE_DESCOMPACTADO)
    try:
        livro = load_workbook(arquivo, read_only=True, data_only=True, keep_links=False)
        try:
            aba, linhas, avisos = _ler_vidas(livro)
        finally:
            livro.close()
    except PlanilhaInvalidaError:
        raise
    except Exception as erro:  # o openpyxl levanta erros de vários tipos com arquivo estragado
        raise PlanilhaInvalidaError(NAO_LEU) from erro
    if not linhas:
        raise PlanilhaInvalidaError(f'A aba "{aba}" não tem nenhuma vida com nome.')
    return PlanilhaLida(getattr(arquivo, "name", "") or "", aba, linhas, avisos)


def _ler_vidas(livro) -> tuple[str, list[Linha], list[str]]:
    """Lê uma aba só: entre as que têm a linha de títulos, a de data mais recente no nome."""
    abas = []
    for folha in livro.worksheets:
        cabecalho = _cabecalho(list(folha.iter_rows(max_row=10, values_only=True)))
        if cabecalho:
            abas.append((folha, *cabecalho))
    if not abas:
        raise PlanilhaInvalidaError(
            "Nenhuma aba da planilha tem a linha de títulos com Nome, CPF, Valor Salário etc."
        )
    folha, numero_do_cabecalho, colunas = max(
        abas, key=lambda aba: data_no_titulo(aba[0].title) or date.min
    )
    avisos = []
    if len(abas) > 1:
        outras = ", ".join(f'"{aba[0].title}"' for aba in abas if aba[0] is not folha)
        avisos.append(
            f'Só a aba "{folha.title}" foi lida; ficou de fora: {outras}. '
            "Para importar outra, salve-a num arquivo separado."
        )
    linhas = []
    dados = folha.iter_rows(min_row=numero_do_cabecalho + 1, values_only=True)
    for numero, valores in enumerate(dados, start=numero_do_cabecalho + 1):
        lidos = {nome: _texto(valores[i], nome) for i, nome in colunas.items() if i < len(valores)}
        nome = chave(lidos.get("nome"))
        if not nome or nome.split()[0] in ("TOTAL", "TOTAIS"):
            continue
        linhas.append(Linha(numero, lidos))
        if len(linhas) > LIMITE_DE_LINHAS:
            raise PlanilhaInvalidaError(f"A planilha passa de {LIMITE_DE_LINHAS} vidas.")
    return folha.title, linhas, avisos


def empacotar(contrato: Contrato, planilha: PlanilhaLida) -> str:
    """As vidas lidas, assinadas, para a confirmação não depender de reenviar o arquivo."""
    return signing.dumps(
        {"contrato": contrato.pk, **asdict(planilha)}, salt=ASSINATURA, compress=True
    )


def desempacotar(texto: str) -> tuple[int, PlanilhaLida]:
    """O contrato e as vidas; levanta signing.BadSignature se foi alterado ou expirou."""
    dados = signing.loads(texto, salt=ASSINATURA, max_age=VALIDADE_DA_PREVIA.total_seconds())
    contrato = dados.pop("contrato")
    dados["linhas"] = [Linha(**linha) for linha in dados["linhas"]]
    return contrato, PlanilhaLida(**dados)


# ----------------------------------------------------------------------
# Os dados de uma vida
# ----------------------------------------------------------------------


def reais(valor: Decimal | None) -> str:
    return "sem valor" if valor is None else f"R$ {number_format(valor, 2, force_grouping=True)}"


def mesmo_nome(um: str, outro: str) -> bool:
    """O mesmo nome, com ou sem os do meio (ex.: Maria Silva e Maria Aparecida Silva)."""
    um, outro = chave(um).split(), chave(outro).split()
    return um == outro or bool(um and outro and um[0] == outro[0] and um[-1] == outro[-1])


def telefone_da_planilha(ddd: str, fone: str) -> str:
    """DDD e FONE juntos, como (11) 98765-4321."""
    ddd, fone = re.sub(r"\D", "", ddd).lstrip("0"), re.sub(r"\D", "", fone)
    if not fone:
        return ""
    numero = ddd + fone if len(ddd) == 2 and len(fone) in (8, 9) else fone
    if len(numero) in (10, 11):
        return f"({numero[:2]}) {numero[2:-4]}-{numero[-4:]}"
    return numero[:20]


def _codigo(texto: str) -> str:
    """O SUB para comparar: 012 e 12 são a mesma subfatura."""
    codigo = chave(texto)
    return (codigo.lstrip("0") or "0") if codigo.isdigit() else codigo


def _mais_comum(textos) -> tuple[str, int]:
    """O texto que mais aparece e quantos textos diferentes havia."""
    contagem = Counter(texto for texto in textos if texto)
    return (contagem.most_common(1)[0][0] if contagem else ""), len(contagem)


@dataclass
class Dados:
    """Os dados de uma vida, já conferidos (o que veio inválido fica de fora, com aviso)."""

    nome: str
    nascimento: date | None
    cpf: str | None
    salario: Decimal | None
    email: str
    telefone: str
    capital: Decimal | None
    fatura: Decimal | None

    @classmethod
    def da_linha(cls, linha: Linha, avisos: list[str]) -> "Dados":
        return cls(
            nome=linha["nome"],
            nascimento=cls._nascimento(linha["nascimento"], avisos),
            cpf=cls._cpf(linha["cpf"], avisos),
            salario=cls._salario(linha["salario"], avisos),
            email=cls._email(linha["email"], avisos),
            telefone=telefone_da_planilha(linha["ddd"], linha["fone"]),
            capital=cls._valor(linha["capital"]),
            fatura=cls._valor(linha["fatura"]),
        )

    @property
    def cpf_formatado(self) -> str:
        return CPF(self.cpf).formatado() if self.cpf else ""

    @staticmethod
    def _nascimento(texto: str, avisos: list[str]) -> date | None:
        if not texto:
            return None
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                data = datetime.strptime(texto, formato).date()
            except ValueError:
                continue
            if date(1900, 1, 1) <= data <= date.today():
                return data
        avisos.append(f"Nascimento inválido na planilha ({texto}): fica de fora.")
        return None

    @staticmethod
    def _cpf(texto: str, avisos: list[str]) -> str | None:
        if not texto:
            return None
        cpf = CPF(CPF(texto).valor.zfill(11))
        if cpf.valido():
            return cpf.valor
        avisos.append(f"CPF inválido na planilha ({texto}): fica de fora.")
        return None

    @staticmethod
    def _valor(texto: str) -> Decimal | None:
        try:
            valor = Decimal(texto)
        except InvalidOperation:
            return None
        return valor if valor.is_finite() and valor > 0 else None

    @classmethod
    def _salario(cls, texto: str, avisos: list[str]) -> Decimal | None:
        valor = cls._valor(texto)
        if valor is None:
            if texto and not re.fullmatch(r"[0.]+", texto):  # zero é o mesmo que em branco
                avisos.append(f"Salário inválido na planilha ({texto}): fica de fora.")
            return None
        return valor.quantize(CENTAVOS, ROUND_HALF_UP)

    @staticmethod
    def _email(texto: str, avisos: list[str]) -> str:
        if not texto:
            return ""
        try:
            validate_email(texto)
        except ValidationError:
            avisos.append(f"E-mail inválido na planilha ({texto}): fica de fora.")
            return ""
        return texto


# ----------------------------------------------------------------------
# A importação
# ----------------------------------------------------------------------


@dataclass
class Resultado:
    """O que acontece com uma vida da planilha."""

    linha: Linha
    dados: Dados
    avisos: list[str]
    situacao: str = ""
    mensagem: str = ""
    # A pessoa no sistema (a nova só existe de fato depois de confirmar).
    pessoa: Beneficiario | None = None
    pessoa_nova: bool = False
    # Capital e fatura da planilha batem com as regras do contrato.
    confere: bool = True


@dataclass
class Grupo:
    """As vidas de uma subfatura (SUB) da planilha e a filial dela no sistema."""

    codigo: str
    resultados: list[Resultado] = field(default_factory=list)
    local: str = ""
    empresa: str = ""
    unidade: Unidade | None = None
    # A SUB já estava cadastrada nesta filial: não há o que escolher.
    cadastrada: bool = False
    opcoes: list[Unidade] = field(default_factory=list)
    mensagens: list[str] = field(default_factory=list)
    problema: str = ""
    # A importação cria ou muda a subfatura.
    muda: bool = False


@dataclass
class Relatorio:
    grupos: list[Grupo]
    # Vidas do contrato no sistema que não estão na planilha (talvez falte uma exclusão).
    fora_da_planilha: list[Vinculo]
    gravado: bool = False

    @property
    def resultados(self) -> list[Resultado]:
        return [resultado for grupo in self.grupos for resultado in grupo.resultados]

    @cached_property
    def contagem(self) -> Counter:
        return Counter(resultado.situacao for resultado in self.resultados)

    @property
    def nao_conferem(self) -> int:
        return sum(1 for resultado in self.resultados if not resultado.confere)

    @property
    def muda_alguma_coisa(self) -> bool:
        situacoes = (IMPORTAR, ATUALIZAR, FILIAL)
        return any(self.contagem[s] for s in situacoes) or any(g.muda for g in self.grupos)

    @property
    def tem_escolha(self) -> bool:
        return any(grupo.opcoes for grupo in self.grupos)


class FaltaFilialError(Exception):
    """Pessoa nova numa subfatura sem filial: a escolha fica para a prévia."""


class ImportacaoDeVidas:
    """Aplica as vidas lidas da planilha num contrato de seguro de vida."""

    def __init__(
        self, contrato: Contrato, planilha: PlanilhaLida, entrada: date, filiais_escolhidas=None
    ):
        self.contrato = contrato
        self.cliente = contrato.cliente
        self.planilha = planilha
        # Quem ainda não está no plano entra nesta data (em geral, a data da planilha).
        self.entrada = entrada
        self.unidades = []
        if self.cliente.eh_empresa:
            self.unidades = list(contrato.unidades.all()) or list(self.cliente.unidades.all())
        por_pk = {unidade.pk: unidade for unidade in self.unidades}
        self.filiais_escolhidas = {
            codigo: por_pk[pk] for codigo, pk in (filiais_escolhidas or {}).items() if pk in por_pk
        }

    def previa(self) -> Relatorio:
        """O que aconteceria com cada vida, sem gravar nada."""
        with transaction.atomic():
            relatorio = self._processar()
            transaction.set_rollback(True)
        return relatorio

    def gravar(self) -> Relatorio:
        """Grava tudo, a não ser que falte escolher a filial de alguma pessoa nova."""
        with transaction.atomic():
            # Duas importações do mesmo cliente ao mesmo tempo (ex.: clique duplo) vão em fila.
            Cliente.objects.select_for_update().get(pk=self.cliente.pk)
            relatorio = self._processar()
            relatorio.gravado = not relatorio.contagem[FILIAL]
            if not relatorio.gravado:
                transaction.set_rollback(True)
        return relatorio

    def _processar(self) -> Relatorio:
        self._carregar_pessoas()
        self.vistas: dict[str, int] = {}
        self.presentes: set[int] = set()
        grupos = self._grupos()
        for grupo in grupos:
            self._salvar_subfatura(grupo)
            for resultado in grupo.resultados:
                self._processar_linha(resultado, grupo)
        return Relatorio(grupos, self._fora_da_planilha())

    def _carregar_pessoas(self) -> None:
        """As pessoas do cliente, por CPF e (os titulares) por nome."""
        self.por_cpf: dict[str, Beneficiario] = {}
        self.titulares_por_nome: dict[str, list[Beneficiario]] = defaultdict(list)
        pessoas = Beneficiario.objects.filter(cliente=self.cliente).select_related(
            "titular", "unidade"
        )
        for pessoa in pessoas:
            self._guardar(pessoa)

    def _guardar(self, pessoa: Beneficiario) -> None:
        if pessoa.cpf:
            self.por_cpf[pessoa.cpf] = pessoa
        if pessoa.eh_titular:
            self.titulares_por_nome[chave(pessoa.nome)].append(pessoa)

    # ------------------------------------------------------------------
    # Subfaturas
    # ------------------------------------------------------------------

    def _grupos(self) -> list[Grupo]:
        """As vidas por subfatura, na ordem da planilha do escritório (SUB crescente)."""
        grupos: dict[str, Grupo] = {}
        for linha in self.planilha.linhas:
            codigo = linha["sub"] if self.cliente.eh_empresa else ""
            grupo = grupos.setdefault(_codigo(codigo), Grupo(codigo))
            avisos: list[str] = []
            grupo.resultados.append(Resultado(linha, Dados.da_linha(linha, avisos), avisos))
        for grupo in grupos.values():
            grupo.local, locais = _mais_comum(r.linha["local"] for r in grupo.resultados)
            grupo.empresa, _ = _mais_comum(r.linha["empresa"] for r in grupo.resultados)
            if locais > 1:
                grupo.mensagens.append(
                    f"Há mais de um Local nesta subfatura: fica o {grupo.local}."
                )
        ordenados = sorted(
            grupos.values(), key=lambda g: (not g.codigo, ordem_do_codigo(_codigo(g.codigo)))
        )
        if self.cliente.eh_empresa:
            self._escolher_filiais(ordenados)
        return ordenados

    def _escolher_filiais(self, grupos: list[Grupo]) -> None:
        """A filial de cada SUB: a já cadastrada, a escolhida na prévia ou a mais provável."""
        ocupadas: dict[int, str] = {}
        subfaturas = list(self.contrato.subfaturas.select_related("unidade"))
        for grupo in grupos:
            cadastrada = next(
                (s for s in subfaturas if mesmo_codigo(s.identificador, grupo.codigo)), None
            )
            if grupo.codigo and cadastrada:
                grupo.unidade, grupo.cadastrada = cadastrada.unidade, True
                ocupadas[cadastrada.unidade_id] = grupo.codigo
        pendentes = [grupo for grupo in grupos if not grupo.cadastrada]
        for grupo in pendentes:
            grupo.opcoes = self.unidades
            if escolhida := self.filiais_escolhidas.get(grupo.codigo):
                self._ocupar(grupo, escolhida, ocupadas)
        for grupo in [g for g in pendentes if g.unidade is None and not g.problema]:
            if provavel := self._filial_provavel(grupo, ocupadas):
                self._ocupar(grupo, provavel, ocupadas)

    @staticmethod
    def _ocupar(grupo: Grupo, unidade: Unidade, ocupadas: dict[int, str]) -> None:
        """Cada filial tem uma subfatura só no contrato (as vidas sem SUB não contam)."""
        if not grupo.codigo:
            grupo.unidade = unidade
        elif unidade.pk in ocupadas:
            grupo.problema = f"A filial {unidade} já está na subfatura {ocupadas[unidade.pk]}."
        else:
            grupo.unidade = unidade
            ocupadas[unidade.pk] = grupo.codigo

    def _filial_provavel(self, grupo: Grupo, ocupadas: dict[int, str]) -> Unidade | None:
        """A filial das pessoas da SUB já cadastradas, a do nome da empresa ou a única."""
        livres = [u for u in self.unidades if not grupo.codigo or u.pk not in ocupadas]
        conhecidas = [self._conhecida(resultado.dados) for resultado in grupo.resultados]
        votos = Counter(pessoa.unidade_id for pessoa in conhecidas if pessoa and pessoa.unidade_id)
        mais_votadas = votos.most_common(2)
        if mais_votadas and (len(mais_votadas) == 1 or mais_votadas[0][1] > mais_votadas[1][1]):
            return next((u for u in livres if u.pk == mais_votadas[0][0]), None)
        if empresa := chave(grupo.empresa):
            nomes = [(u, chave(u.razao_social), chave(u.nome_fantasia)) for u in livres]
            pelo_nome = [u for u, razao, fantasia in nomes if empresa in (razao, fantasia)]
            if len(pelo_nome) == 1:
                return pelo_nome[0]
        return livres[0] if len(self.unidades) == 1 and livres else None

    def _salvar_subfatura(self, grupo: Grupo) -> None:
        """Guarda Local e SUB da filial no contrato (o boleto separado fica como está)."""
        if not (grupo.codigo and grupo.unidade) or grupo.problema:
            return
        subfatura = Subfatura.objects.filter(contrato=self.contrato, unidade=grupo.unidade).first()
        if subfatura is None:
            subfatura = Subfatura(contrato=self.contrato, unidade=grupo.unidade)
            grupo.mensagens.append("Subfatura nova nesta filial.")
        elif not mesmo_codigo(subfatura.identificador, grupo.codigo):
            antes = subfatura.identificador or "vazia"
            grupo.mensagens.append(f"A subfatura desta filial muda de {antes} para {grupo.codigo}.")
        if grupo.local and subfatura.pk and subfatura.local != grupo.local:
            grupo.mensagens.append(
                f"Local muda de {subfatura.local or 'vazio'} para {grupo.local}."
            )
        mudou = not subfatura.pk or not mesmo_codigo(subfatura.identificador, grupo.codigo)
        mudou = mudou or bool(grupo.local and subfatura.local != grupo.local)
        if not mudou:
            return
        subfatura.identificador = grupo.codigo
        subfatura.local = grupo.local or subfatura.local
        try:
            with transaction.atomic():
                subfatura.full_clean()
                subfatura.save()
        except ValidationError as erro:
            grupo.mensagens.clear()
            grupo.problema = "Subfatura não gravada: " + " ".join(erro.messages)
            return
        grupo.muda = True

    # ------------------------------------------------------------------
    # Uma vida
    # ------------------------------------------------------------------

    def _processar_linha(self, resultado: Resultado, grupo: Grupo) -> None:
        try:
            with transaction.atomic():
                resultado.situacao, resultado.mensagem = self._importar(resultado, grupo)
        except (FaltaFilialError, ValidationError, DatabaseError) as motivo:
            # A linha foi desfeita no banco: as pessoas em memória voltam a ser as do banco.
            self._carregar_pessoas()
            if resultado.pessoa_nova:
                resultado.pessoa, resultado.pessoa_nova = None, False
            resultado.situacao, resultado.mensagem = self._sem_importar(motivo)

    @staticmethod
    def _sem_importar(motivo: Exception) -> tuple[str, str]:
        if isinstance(motivo, FaltaFilialError):
            return FILIAL, str(motivo)
        if isinstance(motivo, ValidationError):
            return PROBLEMA, " ".join(motivo.messages)
        return PROBLEMA, "Não foi possível gravar esta vida. Confira os dados da pessoa."

    def _importar(self, resultado: Resultado, grupo: Grupo) -> tuple[str, str]:
        dados = resultado.dados
        self._conferir_repetida(resultado)
        pessoa = self._achar(dados, resultado)
        if pessoa is None:
            pessoa = self._nova_pessoa(resultado, grupo)
            mudancas = []
        else:
            mudancas = self._completar_cadastro(pessoa, resultado, grupo)
        vinculo, entrou, no_plano = self._no_plano(pessoa, resultado)
        self._conferir_cobranca(vinculo, resultado)
        if resultado.pessoa_nova:
            return IMPORTAR, f"Pessoa nova, no plano a partir de {self.entrada:%d/%m/%Y}."
        if entrou:
            return IMPORTAR, " ".join([f"Entra no plano em {self.entrada:%d/%m/%Y}.", *mudancas])
        if no_plano or mudancas:
            return ATUALIZAR, " ".join(["Já está no plano.", *no_plano, *mudancas])
        return PULAR, "Já está no plano, sem mudança."

    def _conferir_repetida(self, resultado: Resultado) -> None:
        dados = resultado.dados
        chaves = [f"{chave(dados.nome)} {dados.nascimento or ''}"]
        if dados.cpf:
            chaves.append(dados.cpf)
        repetida = next((self.vistas[c] for c in chaves if c in self.vistas), None)
        for c in chaves:
            self.vistas.setdefault(c, resultado.linha.numero)
        if repetida is not None:
            raise ValidationError(f"Repetida: a mesma pessoa está na linha {repetida}.")

    def _conhecida(self, dados: Dados) -> Beneficiario | None:
        """A pessoa já cadastrada desta vida, se não houver dúvida (para sugerir a filial)."""
        try:
            return self._achar(dados)
        except ValidationError:
            return None

    def _achar(self, dados: Dados, resultado: Resultado | None = None) -> Beneficiario | None:
        """A pessoa no sistema: pelo CPF ou, sem ele, pelo nome (e nascimento, se preciso)."""
        pelo_cpf = self.por_cpf.get(dados.cpf) if dados.cpf else None
        if pelo_cpf:
            self._encontrada(pelo_cpf, resultado)
            if not pelo_cpf.eh_titular:
                raise ValidationError(
                    f"No sistema, este CPF é de {pelo_cpf.nome}, dependente de "
                    f"{pelo_cpf.titular.nome}."
                )
            if not mesmo_nome(pelo_cpf.nome, dados.nome):
                raise ValidationError(
                    f"No sistema, este CPF é de {pelo_cpf.nome}. Confira o nome ou o CPF."
                )
            return pelo_cpf
        achadas = self.titulares_por_nome.get(chave(dados.nome), [])
        if len(achadas) > 1:
            nascimento = dados.nascimento
            achadas = [p for p in achadas if nascimento and p.data_nascimento == nascimento]
            if len(achadas) != 1:
                raise ValidationError(
                    "Há mais de um titular com este nome no sistema: informe o CPF na planilha."
                )
        if not achadas:
            return None
        pessoa = achadas[0]
        self._encontrada(pessoa, resultado)
        if dados.cpf and pessoa.cpf and pessoa.cpf != dados.cpf:
            raise ValidationError(
                f"No sistema, {pessoa.nome} tem outro CPF ({pessoa.cpf_formatado})."
            )
        nascimento = pessoa.data_nascimento
        if dados.nascimento and nascimento and nascimento != dados.nascimento:
            raise ValidationError(f"No sistema, {pessoa.nome} nasceu em {nascimento:%d/%m/%Y}.")
        return pessoa

    def _encontrada(self, pessoa: Beneficiario, resultado: Resultado | None) -> None:
        if resultado is not None:
            resultado.pessoa = pessoa
            self.presentes.add(pessoa.pk)

    def _nova_pessoa(self, resultado: Resultado, grupo: Grupo) -> Beneficiario:
        dados = resultado.dados
        unidade = None
        if self.cliente.eh_empresa:
            unidade = grupo.unidade
            if unidade is None:
                raise FaltaFilialError("Pessoa nova: escolha a filial acima.")
        pessoa = Beneficiario(
            cliente=self.cliente,
            unidade=unidade,
            nome=dados.nome,
            cpf=dados.cpf,
            data_nascimento=dados.nascimento,
            email=dados.email,
            telefone=dados.telefone,
        )
        pessoa.full_clean()
        pessoa.save()
        self._guardar(pessoa)
        self.presentes.add(pessoa.pk)
        resultado.pessoa, resultado.pessoa_nova = pessoa, True
        return pessoa

    def _completar_cadastro(
        self, pessoa: Beneficiario, resultado: Resultado, grupo: Grupo
    ) -> list[str]:
        """Preenche o que falta no cadastro; o que já está preenchido fica como está."""
        dados = resultado.dados
        preenchidos = []
        campos = (
            ("data_nascimento", dados.nascimento, "nascimento"),
            ("cpf", dados.cpf, "CPF"),
            ("email", dados.email, "e-mail"),
            ("telefone", dados.telefone, "telefone"),
        )
        for campo, valor, rotulo in campos:
            atual = getattr(pessoa, campo)
            if valor and not atual:
                setattr(pessoa, campo, valor)
                preenchidos.append((campo, rotulo))
            elif valor and campo in ("email", "telefone") and not self._igual(campo, atual, valor):
                resultado.avisos.append(f"O {rotulo} da planilha é outro; fica o do cadastro.")
        if grupo.codigo and grupo.unidade and pessoa.unidade_id != grupo.unidade.pk:
            resultado.avisos.append(
                f"No cadastro, a pessoa é da filial {pessoa.unidade}, e não da {grupo.unidade}."
            )
        if not preenchidos:
            return []
        pessoa.full_clean()
        pessoa.save(update_fields=[campo for campo, _ in preenchidos] + ["atualizado_em"])
        if pessoa.cpf:
            self.por_cpf[pessoa.cpf] = pessoa
        rotulos = [rotulo for _, rotulo in preenchidos]
        texto = ", ".join(rotulos[:-1]) + (" e " if len(rotulos) > 1 else "") + rotulos[-1]
        return [f"Cadastro: {texto} {'preenchidos' if len(rotulos) > 1 else 'preenchido'}."]

    @staticmethod
    def _igual(campo: str, atual: str, valor: str) -> bool:
        if campo == "telefone":
            return re.sub(r"\D", "", atual) == re.sub(r"\D", "", valor)
        return atual.strip().lower() == valor.strip().lower()

    def _no_plano(
        self, pessoa: Beneficiario, resultado: Resultado
    ) -> tuple[Vinculo, bool, list[str]]:
        """O plano de vida da pessoa, com o salário da planilha: (plano, entrou, mudanças)."""
        salario = resultado.dados.salario if self.contrato.usa_salario else None
        planos = pessoa.vinculos.filter(contrato=self.contrato)
        atual = (
            planos.filter(Q(data_exclusao__isnull=True) | Q(data_exclusao__gt=self.entrada))
            .order_by("-data_inclusao")
            .first()
        )
        if atual is None:
            saiu = planos.filter(data_exclusao__lte=self.entrada).order_by("-data_exclusao").first()
            if saiu:
                resultado.avisos.append(
                    f"Tinha saído deste plano em {saiu.data_exclusao:%d/%m/%Y}."
                )
            novo = Vinculo(
                beneficiario=pessoa,
                contrato=self.contrato,
                data_inclusao=self.entrada,
                salario=salario,
            )
            novo.full_clean()
            novo.save()
            return novo, True, []
        if salario is None:
            if atual.salario is not None and self.contrato.usa_salario:
                resultado.avisos.append(f"Fica o salário do cadastro, {reais(atual.salario)}.")
            return atual, False, []
        if salario == atual.salario:
            return atual, False, []
        if atual.salario is None:
            mudanca = f"Salário preenchido: {reais(salario)}."
        else:
            mudanca = f"Salário de {reais(atual.salario)} para {reais(salario)}."
        atual.salario = salario
        atual.save(update_fields=["salario", "atualizado_em"])
        return atual, False, [mudanca]

    def _conferir_cobranca(self, vinculo: Vinculo, resultado: Resultado) -> None:
        """Compara o capital e a fatura da planilha com a conta pelas regras do contrato."""
        capital = self.contrato.capital_de(vinculo.salario)
        fatura = self.contrato.fatura_de(capital)
        diferencas = []
        for rotulo, na_planilha, no_contrato in (
            ("capital", resultado.dados.capital, capital),
            ("fatura", resultado.dados.fatura, fatura),
        ):
            if na_planilha is None or no_contrato is None:
                continue
            if abs(na_planilha - no_contrato) > TOLERANCIA:
                diferencas.append(
                    f"{rotulo} {reais(na_planilha)} na planilha e "
                    f"{reais(no_contrato.quantize(CENTAVOS))} pelo contrato"
                )
        if diferencas:
            resultado.confere = False
            resultado.avisos.append("Não bate: " + "; ".join(diferencas) + ".")

    def _fora_da_planilha(self) -> list[Vinculo]:
        return list(
            Vinculo.objects.ativos(self.entrada)
            .filter(contrato=self.contrato, beneficiario__titular__isnull=True)
            .exclude(beneficiario__in=self.presentes)
            .select_related("beneficiario__unidade")
            .order_by("beneficiario__nome")
        )
