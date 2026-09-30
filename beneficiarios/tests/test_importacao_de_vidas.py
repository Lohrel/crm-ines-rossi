"""Importação da planilha do seguro de vida (planilhas fictícias, no layout da seguradora)."""

import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook

from apolices.models import Subfatura
from beneficiarios import importacao_de_vidas
from beneficiarios.importacao_de_vidas import (
    ATUALIZAR,
    FILIAL,
    IMPORTAR,
    PROBLEMA,
    PULAR,
    ImportacaoDeVidas,
    PlanilhaInvalidaError,
    empacotar,
    ler_planilha,
    telefone_da_planilha,
    texto_do_valor,
)
from beneficiarios.models import Beneficiario, Vinculo
from beneficiarios.seguro_de_vida import TIPO_XLSX, PlanilhaDoSeguroDeVida
from beneficiarios.tests.fixtures import contrato
from beneficiarios.tests.test_seguro_de_vida import na_vida
from clientes.tests.fixtures import CPF_1, CPF_2

pytestmark = pytest.mark.django_db

OUTUBRO = date(2026, 10, 1)
ABA = "seguro de vida 01.10.2026"
# CPFs fictícios, com dígitos verificadores válidos.
CPF_COM_ZERO = "01234567890"
CPF_3 = "39813614668"
FILIAL_LTDA = "Supermercados Exemplo Filial Ltda"
# Títulos nas posições da planilha da seguradora (as outras colunas ficam vazias).
TITULOS = {
    "A": "Local",
    "B": "IDENTIFICADOR DE SUB",
    "C": "Empresa",
    "E": "Nome",
    "H": "Nascimento",
    "N": "CPF",
    "O": "Valor Salário",
    "P": "E-MAIL",
    "Q": "DDD",
    "R": "FONE",
    "S": "24x",
    "T": "Capital Segurado",
    "U": "Fatura",
}
COLUNA = {
    "local": "A",
    "sub": "B",
    "empresa": "C",
    "nascimento": "H",
    "cpf": "N",
    "salario": "O",
    "email": "P",
    "ddd": "Q",
    "fone": "R",
    "capital": "T",
    "fatura": "U",
}


def na_planilha(nome, sub=12, local=3, empresa=FILIAL_LTDA, **dados) -> dict:
    """Uma vida, como linha da planilha (coluna → valor)."""
    colunas = {"E": nome, "A": local, "B": sub, "C": empresa}
    colunas.update({COLUNA[campo]: valor for campo, valor in dados.items()})
    return colunas


def planilha(*vidas, aba=ABA, livro=None) -> BytesIO:
    """Títulos na linha 1, uma vida por linha e o total da fatura no fim, como no escritório."""
    livro = livro or Workbook()
    folha = livro.active
    if folha["A1"].value is not None:  # a aba de outra chamada: esta vai numa aba nova
        folha = livro.create_sheet()
    folha.title = aba
    for letra, titulo in TITULOS.items():
        folha[f"{letra}1"] = titulo
    for numero, colunas in enumerate(vidas, start=2):
        for letra, valor in colunas.items():
            folha[f"{letra}{numero}"] = valor
    folha[f"U{len(vidas) + 2}"] = 123.45
    arquivo = BytesIO()
    livro.save(arquivo)
    arquivo.seek(0)
    return arquivo


def importar(contrato, arquivo, entrada=OUTUBRO, filiais=None):
    return ImportacaoDeVidas(contrato, ler_planilha(arquivo), entrada, filiais).gravar()


def por_linha(relatorio) -> dict[int, tuple[str, str]]:
    return {r.linha.numero: (r.situacao, r.mensagem) for r in relatorio.resultados}


# ------------------------------------------------------------------
# Leitura
# ------------------------------------------------------------------


def test_le_a_planilha_no_layout_da_seguradora():
    arquivo = planilha(
        na_planilha(
            "Ana Titular",
            nascimento="05/03/1985",
            cpf=int(CPF_1),
            salario=2500.5,
            capital=60012.0,
            fatura=21.0042,
        ),
        na_planilha(
            "Bia Nova",
            sub=10,
            local=1,
            nascimento=datetime(1990, 7, 1),
            cpf=int(CPF_COM_ZERO),
            salario=3000,
            email="bia@exemplo.test",
            ddd=11,
            fone=987654321,
        ),
    )
    lida = ler_planilha(arquivo)
    assert (lida.aba, lida.data, lida.avisos) == (ABA, OUTUBRO, [])
    # A linha do total, sem nome, fica de fora.
    assert [linha.numero for linha in lida.linhas] == [2, 3]
    ana, bia = lida.linhas
    assert (ana["sub"], ana["local"], ana["nascimento"], ana["cpf"]) == (
        "12",
        "3",
        "05/03/1985",
        CPF_1,
    )
    assert (ana["salario"], ana["capital"], ana["fatura"]) == ("2500.5", "60012", "21.0042")
    # CPF guardado como número volta com o zero da frente.
    assert (bia["nascimento"], bia["cpf"], bia["salario"]) == ("01/07/1990", CPF_COM_ZERO, "3000")
    assert (bia["ddd"], bia["fone"], bia["email"]) == ("11", "987654321", "bia@exemplo.test")


def test_le_so_a_aba_mais_recente():
    livro = Workbook()
    planilha(na_planilha("Ana Antiga"), aba="seguro de vida 01.09.2026", livro=livro)
    arquivo = planilha(na_planilha("Bia Atual"), livro=livro)
    lida = ler_planilha(arquivo)
    assert [linha["nome"] for linha in lida.linhas] == ["Bia Atual"]
    assert lida.avisos == [
        f'Só a aba "{ABA}" foi lida; ficou de fora: "seguro de vida 01.09.2026". '
        "Para importar outra, salve-a num arquivo separado."
    ]


def test_arquivo_que_nao_e_planilha():
    with pytest.raises(PlanilhaInvalidaError, match="Não foi possível ler o arquivo"):
        ler_planilha(BytesIO(b"isto nao e uma planilha"))


def test_planilha_sem_a_linha_de_titulos():
    livro = Workbook()
    livro.active.append(["Relatório", "qualquer"])
    arquivo = BytesIO()
    livro.save(arquivo)
    with pytest.raises(PlanilhaInvalidaError, match="Nenhuma aba da planilha tem a linha"):
        ler_planilha(arquivo)


def test_planilha_sem_nomes():
    with pytest.raises(PlanilhaInvalidaError, match="não tem nenhuma vida com nome"):
        ler_planilha(planilha(na_planilha(None, salario=1000)))


def test_planilha_grande_demais(monkeypatch):
    monkeypatch.setattr(importacao_de_vidas, "LIMITE_DESCOMPACTADO", 100)
    with pytest.raises(PlanilhaInvalidaError, match="grande demais"):
        ler_planilha(planilha(na_planilha("Ana Nova")))


def test_limite_de_vidas(monkeypatch):
    monkeypatch.setattr(importacao_de_vidas, "LIMITE_DE_LINHAS", 2)
    vidas = [na_planilha(nome) for nome in ("Ana Nova", "Bia Nova", "Caio Novo")]
    with pytest.raises(PlanilhaInvalidaError, match="passa de 2 vidas"):
        ler_planilha(planilha(*vidas))


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (2500.5, "2500.5"),
        (3000, "3000"),
        ("2.500,50", "2500.50"),
        ("R$ 1.234.567,89", "1234567.89"),
        ("2500,5", "2500.5"),
        ("2500.50", "2500.50"),
        ("-", ""),
        ("abc", "abc"),
    ],
)
def test_valores_em_reais(valor, esperado):
    assert texto_do_valor(valor) == esperado


@pytest.mark.parametrize(
    ("ddd", "fone", "esperado"),
    [
        ("11", "987654321", "(11) 98765-4321"),
        ("011", "33334444", "(11) 3333-4444"),
        ("", "1133334444", "(11) 3333-4444"),
        ("", "33334444", "33334444"),
        ("11", "", ""),
    ],
)
def test_telefone_da_planilha(ddd, fone, esperado):
    assert telefone_da_planilha(ddd, fone) == esperado


# ------------------------------------------------------------------
# Importação
# ------------------------------------------------------------------


def test_previa_nao_grava_nada(vida, filial, titular):
    lida = ler_planilha(planilha(na_planilha("Ana Titular"), na_planilha("Bia Nova")))
    previa = ImportacaoDeVidas(vida, lida, OUTUBRO).previa()
    assert previa.contagem == Counter({IMPORTAR: 2})
    assert previa.muda_alguma_coisa
    assert Beneficiario.objects.count() == 1
    assert not Vinculo.objects.exists() and not Subfatura.objects.exists()


def test_importa_pessoas_vidas_e_subfaturas(vida, matriz, filial, titular):
    relatorio = importar(
        vida,
        planilha(
            # Ana já está cadastrada na filial, mas ainda não no seguro de vida.
            na_planilha("Ana Titular", salario=2500.5, capital=60012.0, fatura=21.0042),
            na_planilha(
                "Bia Nova",
                nascimento="01/07/1990",
                cpf=CPF_3,
                salario=3000,
                email="bia@exemplo.test",
                ddd=11,
                fone=987654321,
            ),
            na_planilha("Caio Novo", sub=10, local=1, empresa="Supermercados Exemplo Ltda"),
        ),
    )
    assert por_linha(relatorio) == {
        2: (IMPORTAR, "Entra no plano em 01/10/2026."),
        3: (IMPORTAR, "Pessoa nova, no plano a partir de 01/10/2026."),
        4: (IMPORTAR, "Pessoa nova, no plano a partir de 01/10/2026."),
    }
    assert relatorio.nao_conferem == 0
    # Na ordem da planilha do escritório: SUB 10 antes da 12.
    sub_10, sub_12 = relatorio.grupos
    assert (sub_10.codigo, sub_10.unidade, sub_10.mensagens) == (
        "10",
        matriz,
        ["Subfatura nova nesta filial."],
    )
    # A filial da SUB 12 veio da Ana, que já estava cadastrada; a da 10, do nome da empresa.
    assert sub_12.unidade == filial
    subfaturas = {s.unidade: (s.identificador, s.local) for s in vida.subfaturas.all()}
    assert subfaturas == {matriz: ("10", "1"), filial: ("12", "3")}
    bia = Beneficiario.objects.get(nome="Bia Nova")
    assert (bia.unidade, bia.cpf, bia.data_nascimento) == (filial, CPF_3, date(1990, 7, 1))
    assert (bia.email, bia.telefone) == ("bia@exemplo.test", "(11) 98765-4321")
    assert Beneficiario.objects.get(nome="Caio Novo").unidade == matriz
    vidas = {v.beneficiario.nome: (v.data_inclusao, v.salario) for v in vida.vinculos.all()}
    assert vidas == {
        "Ana Titular": (OUTUBRO, Decimal("2500.50")),
        "Bia Nova": (OUTUBRO, Decimal("3000.00")),
        "Caio Novo": (OUTUBRO, None),
    }


def test_quem_ja_esta_no_plano_so_tem_o_salario_atualizado(vida, filial, titular):
    na_vida(titular, vida, "2000")
    arquivo = planilha(na_planilha("Ana Titular", salario=2500))
    relatorio = importar(vida, arquivo)
    assert por_linha(relatorio) == {
        2: (ATUALIZAR, "Já está no plano. Salário de R$ 2.000,00 para R$ 2.500,00.")
    }
    vinculo = titular.vinculos.get()
    assert (vinculo.salario, vinculo.data_inclusao) == (Decimal("2500.00"), date(2026, 3, 1))
    # A mesma planilha de novo não muda nada.
    arquivo.seek(0)
    de_novo = importar(vida, arquivo)
    assert por_linha(de_novo) == {2: (PULAR, "Já está no plano, sem mudança.")}
    assert not de_novo.muda_alguma_coisa
    assert de_novo.grupos[0].cadastrada


def test_sem_salario_na_planilha_fica_o_do_cadastro(vida, filial, titular):
    na_vida(titular, vida, "2000")
    resultado = importar(vida, planilha(na_planilha("Ana Titular"))).resultados[0]
    assert resultado.situacao == PULAR
    assert resultado.avisos == ["Fica o salário do cadastro, R$ 2.000,00."]
    assert titular.vinculos.get().salario == Decimal("2000.00")


def test_completa_o_cadastro_sem_apagar_o_que_ja_existe(vida, filial, titular):
    titular.email = "ana@exemplo.test"
    titular.save()
    arquivo = planilha(
        na_planilha(
            "ANA TITULAR",
            nascimento="05/03/1985",
            email="outro@exemplo.test",
            ddd=11,
            fone=33334444,
        )
    )
    resultado = importar(vida, arquivo).resultados[0]
    assert resultado.mensagem == (
        "Entra no plano em 01/10/2026. Cadastro: nascimento e telefone preenchidos."
    )
    assert resultado.avisos == ["O e-mail da planilha é outro; fica o do cadastro."]
    titular.refresh_from_db()
    assert titular.data_nascimento == date(1985, 3, 5)
    assert (titular.telefone, titular.email) == ("(11) 3333-4444", "ana@exemplo.test")
    assert titular.history.first().history_type == "~"


def test_conflitos_ficam_de_fora(vida, filial, titular, conjuge):
    conjuge.cpf = CPF_2
    conjuge.save()
    titular.data_nascimento = date(1985, 3, 5)
    titular.save()
    relatorio = importar(
        vida,
        planilha(
            na_planilha("Outra Pessoa", cpf=CPF_1),
            na_planilha("Bruno Cônjuge", cpf=CPF_2),
            na_planilha("Ana Titular", nascimento="06/03/1985"),
            na_planilha("Dani Nova", salario=1000),
            na_planilha("Dani Nova", salario=1000),
        ),
    )
    assert por_linha(relatorio) == {
        2: (PROBLEMA, "No sistema, este CPF é de Ana Titular. Confira o nome ou o CPF."),
        3: (PROBLEMA, "No sistema, este CPF é de Bruno Cônjuge, dependente de Ana Titular."),
        4: (PROBLEMA, "No sistema, Ana Titular nasceu em 05/03/1985."),
        5: (IMPORTAR, "Pessoa nova, no plano a partir de 01/10/2026."),
        6: (PROBLEMA, "Repetida: a mesma pessoa está na linha 5."),
    }
    # A pessoa do conflito aparece na prévia com o link para o cadastro.
    assert relatorio.resultados[0].pessoa == titular
    assert list(vida.vinculos.values_list("beneficiario__nome", flat=True)) == ["Dani Nova"]


def test_mesmo_cpf_com_nome_abreviado_e_a_mesma_pessoa(vida, filial, titular):
    titular.nome = "Ana Maria Titular"
    titular.save()
    relatorio = importar(vida, planilha(na_planilha("Ana Titular", cpf=CPF_1)))
    assert relatorio.resultados[0].situacao == IMPORTAR
    assert titular.vinculos.filter(contrato=vida).exists()


def test_dois_titulares_com_o_mesmo_nome(vida, empresa, filial, titular):
    Beneficiario.objects.create(
        cliente=empresa, unidade=filial, nome="Ana Titular", data_nascimento=date(1990, 1, 2)
    )
    relatorio = importar(
        vida,
        planilha(
            na_planilha("Ana Titular", nascimento="02/01/1990"),
            na_planilha("Ana Titular", nascimento="03/01/1990"),
        ),
    )
    assert por_linha(relatorio) == {
        2: (IMPORTAR, "Entra no plano em 01/10/2026."),
        3: (
            PROBLEMA,
            "Há mais de um titular com este nome no sistema: informe o CPF na planilha.",
        ),
    }


def test_valores_invalidos_ficam_de_fora_com_aviso(vida, filial):
    arquivo = planilha(
        na_planilha(
            "Ana Nova", nascimento="31/02/1985", cpf="123", salario="abc", email="sem-arroba"
        )
    )
    resultado = importar(vida, arquivo).resultados[0]
    assert resultado.situacao == IMPORTAR
    assert resultado.avisos == [
        "Nascimento inválido na planilha (31/02/1985): fica de fora.",
        "CPF inválido na planilha (123): fica de fora.",
        "Salário inválido na planilha (abc): fica de fora.",
        "E-mail inválido na planilha (sem-arroba): fica de fora.",
    ]


def test_pessoa_nova_espera_a_escolha_da_filial(vida, matriz, filial):
    arquivo = planilha(na_planilha("Eva Nova", sub=20, empresa="Empresa Sem Cadastro"))
    lida = ler_planilha(arquivo)
    previa = ImportacaoDeVidas(vida, lida, OUTUBRO).previa()
    assert por_linha(previa) == {2: (FILIAL, "Pessoa nova: escolha a filial acima.")}
    assert (previa.grupos[0].unidade, previa.grupos[0].opcoes) == (None, [matriz, filial])
    assert previa.tem_escolha
    # Sem a filial, a confirmação não grava nada.
    assert not ImportacaoDeVidas(vida, lida, OUTUBRO).gravar().gravado
    assert not Beneficiario.objects.exists()
    relatorio = ImportacaoDeVidas(vida, lida, OUTUBRO, {"20": matriz.pk}).gravar()
    assert relatorio.gravado and relatorio.contagem == Counter({IMPORTAR: 1})
    assert Beneficiario.objects.get(nome="Eva Nova").unidade == matriz
    assert vida.subfaturas.get().identificador == "20"


def test_cada_filial_tem_uma_subfatura_so(vida, matriz, filial):
    Subfatura.objects.create(contrato=vida, unidade=matriz, identificador="10", local="1")
    arquivo = planilha(
        na_planilha("Eva Nova", sub="010", local=1),
        na_planilha("Gil Novo", sub=20, empresa="Empresa Sem Cadastro"),
    )
    relatorio = ImportacaoDeVidas(vida, ler_planilha(arquivo), OUTUBRO, {"20": matriz.pk}).previa()
    sub_10, sub_20 = relatorio.grupos
    # 010 e 10 são a mesma subfatura, já cadastrada na matriz.
    assert (sub_10.unidade, sub_10.cadastrada, sub_10.opcoes) == (matriz, True, [])
    assert sub_20.problema == "A filial Supermercados Exemplo Ltda já está na subfatura 010."
    assert por_linha(relatorio)[3][0] == FILIAL


def test_confere_capital_e_fatura_com_o_contrato(vida, filial):
    relatorio = importar(
        vida,
        planilha(
            na_planilha("Ana Certa", salario=2500.5, capital=60012.0, fatura=21.0042),
            na_planilha("Bia Teto", salario=20000, capital=300000, fatura=105),
            na_planilha("Caio Errado", salario=3000, capital=144000, fatura=50.4),
        ),
    )
    assert relatorio.nao_conferem == 1
    caio = relatorio.resultados[2]
    assert caio.avisos == [
        "Não bate: capital R$ 144.000,00 na planilha e R$ 72.000,00 pelo contrato; "
        "fatura R$ 50,40 na planilha e R$ 25,20 pelo contrato."
    ]


def test_capital_fixo_nao_guarda_salario(vida, filial):
    vida.capital_multiplo = None
    vida.save()
    arquivo = planilha(na_planilha("Ana Nova", salario=2500, capital=300000, fatura=105))
    relatorio = importar(vida, arquivo)
    assert (relatorio.contagem, relatorio.nao_conferem) == (Counter({IMPORTAR: 1}), 0)
    assert Vinculo.objects.get(beneficiario__nome="Ana Nova").salario is None


def test_quem_tinha_saido_volta_ao_plano(vida, filial, titular):
    na_vida(titular, vida, "2000", exclusao=date(2026, 6, 1))
    resultado = importar(vida, planilha(na_planilha("Ana Titular", salario=2100))).resultados[0]
    assert (resultado.situacao, resultado.avisos) == (
        IMPORTAR,
        ["Tinha saído deste plano em 01/06/2026."],
    )
    assert titular.vinculos.get(data_exclusao__isnull=True).data_inclusao == OUTUBRO


def test_pessoa_de_outra_filial_fica_onde_esta(vida, matriz, filial, titular):
    Subfatura.objects.create(contrato=vida, unidade=matriz, identificador="10")
    resultado = importar(vida, planilha(na_planilha("Ana Titular", sub=10))).resultados[0]
    assert resultado.avisos == [
        "No cadastro, a pessoa é da filial Exemplo Centro, e não da Supermercados Exemplo Ltda."
    ]
    titular.refresh_from_db()
    assert titular.unidade == filial


def test_lista_as_vidas_do_sistema_fora_da_planilha(vida, empresa, filial, titular):
    na_vida(titular, vida, "2000")
    saiu = Beneficiario.objects.create(cliente=empresa, unidade=filial, nome="Gil Saiu")
    na_vida(saiu, vida, "3000", exclusao=OUTUBRO)
    relatorio = importar(vida, planilha(na_planilha("Bia Nova")))
    assert [v.beneficiario for v in relatorio.fora_da_planilha] == [titular]


def test_cliente_pessoa_fisica(pessoa, seguradora):
    individual = contrato(pessoa, seguradora, "Vida", "VI-1", [])
    individual.capital_multiplo = 24
    individual.capital_maximo = Decimal("300000")
    individual.save()
    relatorio = importar(individual, planilha(na_planilha("João Exemplo", sub=5, salario=4000)))
    assert relatorio.contagem == Counter({IMPORTAR: 1})
    assert (relatorio.grupos[0].codigo, relatorio.grupos[0].opcoes) == ("", [])
    joao = Beneficiario.objects.get(cliente=pessoa)
    assert joao.unidade is None
    assert joao.vinculos.get().salario == Decimal("4000.00")
    assert not Subfatura.objects.exists()


def test_a_planilha_baixada_do_sistema_volta_sem_mudanca(vida, empresa, matriz, filial, titular):
    Subfatura.objects.create(contrato=vida, unidade=matriz, local="1", identificador="10")
    Subfatura.objects.create(contrato=vida, unidade=filial, local="3", identificador="2")
    titular.data_nascimento = date(1990, 4, 12)
    titular.email = "ana@exemplo.test"
    titular.telefone = "(11) 98765-4321"
    titular.save()
    na_vida(titular, vida, "2500")
    na_vida(
        Beneficiario.objects.create(cliente=empresa, unidade=matriz, nome="Bia"), vida, "1000.50"
    )
    na_vida(Beneficiario.objects.create(cliente=empresa, unidade=matriz, nome="Caio"), vida, None)
    exportada = PlanilhaDoSeguroDeVida(vida, OUTUBRO)
    lida = ler_planilha(BytesIO(exportada.gerar()))
    assert lida.data == OUTUBRO
    assert [linha["nome"] for linha in lida.linhas] == ["Ana Titular", "Bia", "Caio"]
    relatorio = ImportacaoDeVidas(vida, lida, OUTUBRO).gravar()
    assert relatorio.contagem == Counter({PULAR: 3})
    assert not relatorio.muda_alguma_coisa and relatorio.fora_da_planilha == []
    assert [r.avisos for r in relatorio.resultados] == [[], [], []]


# ------------------------------------------------------------------
# Tela
# ------------------------------------------------------------------


def url(contrato) -> str:
    return reverse("beneficiarios:importar_vidas", args=[contrato.pk])


def enviar(cliente_http, contrato, arquivo):
    anexo = SimpleUploadedFile("vidas.xlsx", arquivo.getvalue(), content_type=TIPO_XLSX)
    return cliente_http.post(url(contrato), {"arquivo": anexo})


def dados_da_previa(resposta) -> str:
    return re.search(r'name="dados" value="([^"]+)"', resposta.content.decode()).group(1)


def test_importar_pela_tela(logado, vida, matriz, filial):
    assert "Enviar a planilha atual" in logado.get(url(vida)).content.decode()
    previa = enviar(logado, vida, planilha(na_planilha("Ana Nova", salario=2500)))
    pagina = previa.content.decode()
    assert "Prévia da importação" in pagina
    assert "Pessoa nova, no plano a partir de 01/10/2026." in pagina
    # A filial sugerida (pelo nome da empresa) já vem escolhida.
    assert f'<option value="{filial.pk}" selected>' in pagina
    assert not Beneficiario.objects.exists()
    resposta = logado.post(
        url(vida),
        {
            "dados": dados_da_previa(previa),
            "entrada": "15/10/2026",
            "filial-12": filial.pk,
            "acao": "confirmar",
        },
    )
    assert "Planilha importada" in resposta.content.decode()
    vinculo = Vinculo.objects.get(beneficiario__nome="Ana Nova")
    assert (vinculo.data_inclusao, vinculo.salario) == (date(2026, 10, 15), Decimal("2500.00"))


def test_confirmar_sem_a_filial_nao_grava(logado, vida, matriz, filial):
    arquivo = planilha(na_planilha("Eva Nova", sub=20, empresa="Empresa Sem Cadastro"))
    previa = enviar(logado, vida, arquivo)
    dados = {"dados": dados_da_previa(previa), "entrada": "01/10/2026", "acao": "confirmar"}
    pagina = logado.post(url(vida), dados).content.decode()
    assert "Escolha a filial das pessoas novas antes de confirmar." in pagina
    assert not Beneficiario.objects.exists()


def test_data_de_entrada_invalida_nao_grava(logado, vida, filial):
    previa = enviar(logado, vida, planilha(na_planilha("Ana Nova")))
    dados = {"dados": dados_da_previa(previa), "entrada": "31/02/2026", "acao": "confirmar"}
    pagina = logado.post(url(vida), dados).content.decode()
    assert "Prévia da importação" in pagina and "is-invalid" in pagina
    assert not Beneficiario.objects.exists()


def test_previa_vencida_ou_alterada(logado, vida):
    resposta = logado.post(url(vida), {"dados": "alterada", "acao": "confirmar"}, follow=True)
    assert resposta.redirect_chain == [(url(vida), 302)]
    assert "A prévia expirou. Envie a planilha de novo." in resposta.content.decode()


def test_previa_de_outro_contrato(logado, vida, empresa, seguradora, matriz):
    outro = contrato(empresa, seguradora, "Vida", "V-2", [matriz])
    outro.capital_maximo = Decimal("50000")
    outro.save()
    dados = empacotar(vida, ler_planilha(planilha(na_planilha("Ana Nova"))))
    assert logado.post(url(outro), {"dados": dados}).status_code == 404


def test_arquivo_invalido_mostra_o_erro(logado, vida):
    pagina = enviar(logado, vida, BytesIO(b"nada")).content.decode()
    assert "Não foi possível ler o arquivo." in pagina


def test_contrato_sem_capital_pede_a_configuracao(logado, saude):
    assert "Antes de importar, configure o capital" in logado.get(url(saude)).content.decode()
    assert logado.post(url(saude), {}).status_code == 302


def test_importar_pede_login(client, vida):
    assert client.get(url(vida)).status_code == 302


def test_atalho_na_planilha_do_seguro_de_vida(logado, vida):
    pagina = logado.get(reverse("beneficiarios:seguro_de_vida", args=[vida.pk])).content.decode()
    assert url(vida) in pagina
