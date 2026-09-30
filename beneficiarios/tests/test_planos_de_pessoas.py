"""Auto e imóvel guardam veículos e imóveis, não pessoas: ficam fora das vidas e dos planos."""

import re
from datetime import date
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apolices.models import Contrato
from beneficiarios.tests.fixtures import contrato, incluir
from calendario.agenda import Agenda
from contas.painel import Painel

pytestmark = pytest.mark.django_db


@pytest.fixture
def frota(empresa, seguradora, matriz, filial):
    return contrato(empresa, seguradora, "Auto", "FROTA-1", [matriz, filial])


def test_so_os_planos_de_pessoas(saude, dental, frota, empresa, seguradora, matriz):
    casa = contrato(empresa, seguradora, "Imóvel", "CASA-1", [matriz])
    assert set(Contrato.objects.de_pessoas()) == {saude, dental}
    assert casa not in Contrato.objects.de_pessoas()


def test_quadro_de_vidas_sem_auto(logado, empresa, titular, saude, frota):
    incluir(titular, saude)
    pagina = logado.get(reverse("beneficiarios:lista", args=[empresa.pk])).content.decode()
    assert "S-1" in pagina and "FROTA-1" not in pagina


def test_plano_da_pessoa_sem_auto(logado, titular, saude, frota):
    pagina = logado.get(reverse("beneficiarios:novo_plano", args=[titular.pk])).content.decode()
    assert "S-1" in pagina and "FROTA-1" not in pagina


def test_inclusao_sem_auto(logado, titular, saude, frota):
    url = reverse("movimentacoes:nova", args=[titular.pk, "inclusao"])
    pagina = logado.get(url).content.decode()
    assert "S-1" in pagina and "FROTA-1" not in pagina


def test_planilha_mensal_sem_aba_de_auto(logado, empresa, saude, frota):
    url = reverse("movimentacoes:planilha")
    resposta = logado.get(url, {"cliente": empresa.pk, "mes": "09/2026"})
    livro = load_workbook(BytesIO(b"".join(resposta.streaming_content)))
    assert livro.sheetnames == ["SAÚDE"]


def test_relatorios_de_vidas_sem_auto(logado, saude, frota):
    pagina = logado.get(reverse("relatorios:vidas")).content.decode()
    assert "S-1" in pagina and "FROTA-1" not in pagina


def test_relatorio_de_apolices_sem_vidas_no_auto(logado, titular, saude, frota):
    incluir(titular, saude)
    url = reverse("relatorios:apolices")
    pagina = logado.get(url).content.decode()
    linha_da_frota = re.search(r"<tr[^>]*>(?:(?!</tr>).)*FROTA-1.*?</tr>", pagina, re.S).group(0)
    assert '<td class="text-end">—</td>' in linha_da_frota
    linha_da_saude = re.search(r"<tr[^>]*>(?:(?!</tr>).)*>S-1<.*?</tr>", pagina, re.S).group(0)
    assert '<td class="text-end">1</td>' in linha_da_saude

    resposta = logado.get(url, {"formato": "excel"})
    aba = load_workbook(BytesIO(b"".join(resposta.streaming_content))).active
    vidas = {linha[7]: linha[-1] for linha in aba.iter_rows(min_row=3, values_only=True)}
    assert vidas == {"S-1": 1, "FROTA-1": None}


def test_fechamento_das_movimentacoes_sem_auto(usuario, saude, frota):
    Contrato.objects.filter(pk__in=[saude.pk, frota.pk]).update(dia_fechamento=25)
    prazos = Painel(usuario, hoje=date(2026, 9, 20)).prazos
    assert [(p["titulo"], p["contratos"]) for p in prazos] == [("Fechamento das movimentações", 1)]
    datas = Agenda(date(2026, 9, 1), date(2026, 9, 30)).por_dia()
    assert [(item.titulo, item.detalhe) for item in datas[date(2026, 9, 25)]] == [
        ("Fechamento das movimentações", "1 contrato fecha neste dia")
    ]
