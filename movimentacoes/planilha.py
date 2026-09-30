"""Planilha mensal de controle das inclusões, no formato que o escritório já usa.

Uma planilha por empresa e por mês. Cada ramo tem uma aba (ex.: SAÚDE, DENTAL)
que junta os planos dele: na Unimed, cada plano (PLENO, FLEX...) é um contrato,
e a coluna PLANO diz em qual a pessoa entrou. As alterações de plano ficam na
aba ALTERAÇÕES DE PLANO. O mês é o da digitação no portal (envio à operadora):
a planilha de setembro traz o que foi digitado em setembro, em qualquer etapa,
e o que foi digitado antes e continua pendente, como o escritório fazia à mão.
"""

import re
from calendar import monthrange
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from io import BytesIO

from django.db.models import Q
from django.db.models.functions import Coalesce
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apolices.models import Contrato
from clientes.models import Cliente
from movimentacoes.models import ETAPAS_EM_ABERTO, Etapa, Movimentacao

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MESES = "janeiro fevereiro março abril maio junho julho agosto setembro outubro novembro dezembro"
MESES = MESES.split()

# A planilha usa só Pendente e Processado; rejeitadas e canceladas aparecem como tal.
STATUS_NA_PLANILHA = {
    Etapa.EM_ANALISE: "Pendente",
    Etapa.APROVADA: "Pendente",
    Etapa.CONCLUIDA: "Processado",
    Etapa.REJEITADA: "Rejeitado",
    Etapa.CANCELADA: "Cancelado",
}
# Na coluna CADASTRO, o titular aparece com a matrícula e o dependente, com esta palavra.
DEPENDENTE = "Dependente"
ALTERACOES = "ALTERAÇÕES DE PLANO"


@dataclass(frozen=True)
class Coluna:
    titulo: str
    largura: int
    data: bool = False


COLUNAS = [
    Coluna("SOLICITANTE", 16),
    Coluna("STATUS", 12),
    Coluna("CADASTRO", 12),
    Coluna("NOME", 34),
    Coluna("CARTEIRINHA", 26),
    Coluna("DATA DIGITAÇÃO", 16, data=True),
    Coluna("INICIO", 12, data=True),
    Coluna("PLANO", 16),
    Coluna("OBSERVAÇÃO", 45),
]
# A aba das alterações não tem carteirinha, como na planilha do escritório.
COLUNAS_DA_ALTERACAO = [coluna for coluna in COLUNAS if coluna.titulo != "CARTEIRINHA"]
# Só nas empresas com mais de uma unidade.
FILIAL = Coluna("FILIAL", 24)


@dataclass
class Aba:
    titulo: str
    cabecalho: str
    colunas: list[Coluna]
    linhas: list[list]


def cadastro(movimentacao: Movimentacao) -> str | int:
    pessoa = movimentacao.beneficiario
    if not pessoa.eh_titular:
        return DEPENDENTE
    # Matrícula só com algarismos vira número, como na planilha do escritório.
    matricula = pessoa.matricula
    return int(matricula) if matricula.isdigit() and not matricula.startswith("0") else matricula


def observacao(movimentacao: Movimentacao) -> str:
    partes = [movimentacao.observacao.strip()]
    if movimentacao.motivo_rejeicao:
        partes.insert(0, f"Rejeitada: {movimentacao.get_motivo_rejeicao_display()}.")
    if movimentacao.eh_alteracao:
        anterior, novo = movimentacao.contrato_anterior, movimentacao.contrato
        partes.insert(0, f"{anterior.nome_do_plano} para {novo.nome_do_plano}.")
    return " ".join(parte for parte in partes if parte)


def nome_de_aba(texto: str) -> str:
    # O Excel não aceita estes caracteres nem mais de 31 letras no nome da aba.
    return re.sub(r"[\[\]:*?/\\]", "-", texto.upper())[:31]


def abas_por_ramo(contratos: list[Contrato]) -> dict[str, list[Contrato]]:
    """Os planos de cada aba: uma por ramo, ou por ramo e seguradora se o ramo tem mais de uma."""
    seguradoras = defaultdict(set)
    for contrato in contratos:
        seguradoras[contrato.ramo_id].add(contrato.seguradora_id)
    abas = defaultdict(list)
    for contrato in contratos:
        titulo = str(contrato.ramo)
        if len(seguradoras[contrato.ramo_id]) > 1:
            titulo = f"{titulo} {contrato.seguradora}"
        abas[nome_de_aba(titulo)].append(contrato)
    return dict(abas)


class PlanilhaDoMes:
    """Inclusões e alterações de uma empresa digitadas num mês, mais as pendentes de antes."""

    def __init__(self, cliente: Cliente, mes: date) -> None:
        self.cliente = cliente
        self.inicio = mes.replace(day=1)
        self.fim = mes.replace(day=monthrange(mes.year, mes.month)[1])
        self.com_filial = cliente.unidades.count() > 1

    @property
    def mes(self) -> str:
        return f"{MESES[self.inicio.month - 1]} de {self.inicio.year}"

    def digitadas(self, tipo: str):
        do_mes = Q(digitada_em__range=(self.inicio, self.fim))
        pendentes_de_antes = Q(digitada_em__lt=self.inicio, etapa__in=ETAPAS_EM_ABERTO)
        return (
            Movimentacao.objects.filter(cliente=self.cliente, tipo=tipo)
            .annotate(digitada_em=Coalesce("enviada_em", "recebida_em"))
            .filter(do_mes | pendentes_de_antes)
            .select_related(
                "solicitante",
                "contrato",
                "contrato_anterior",
                "beneficiario__unidade",
                "beneficiario__titular__unidade",
            )
            .order_by("digitada_em", "beneficiario__nome", "pk")
        )

    def contratos(self) -> list[Contrato]:
        """Os planos vigentes do cliente e qualquer outro com inclusão na planilha."""
        com_inclusao = self.digitadas(Movimentacao.Tipo.INCLUSAO).values("contrato")
        return list(
            Contrato.objects.de_pessoas()
            .filter(
                Q(cliente=self.cliente),
                Q(pk__in=com_inclusao) | Q(situacao__in=Contrato.SITUACOES_VIGENTES),
            )
            .select_related("ramo", "seguradora")
            .order_by("ramo__nome", "seguradora__nome", "produto", "numero")
        )

    def abas(self) -> list[Aba]:
        grupos = abas_por_ramo(self.contratos())
        aba_do_plano = {plano.pk: titulo for titulo, planos in grupos.items() for plano in planos}
        linhas = defaultdict(list)
        for mov in self.digitadas(Movimentacao.Tipo.INCLUSAO):
            linhas[aba_do_plano[mov.contrato_id]].append(self._linha(mov, COLUNAS))
        abas = []
        for titulo, planos in grupos.items():
            seguradoras = ", ".join(dict.fromkeys(str(plano.seguradora) for plano in planos))
            cabecalho = f"{seguradoras} · {planos[0].ramo} · {self.cliente} · {self.mes}"
            abas.append(Aba(titulo, cabecalho, self._colunas(COLUNAS), linhas[titulo]))
        # A aba das alterações aparece quando houve alguma, ou quando há como haver.
        alteracoes = [
            self._linha(mov, COLUNAS_DA_ALTERACAO)
            for mov in self.digitadas(Movimentacao.Tipo.ALTERACAO)
        ]
        if alteracoes or any(len(planos) > 1 for planos in grupos.values()):
            cabecalho = f"Alterações de plano · {self.cliente} · {self.mes}"
            colunas = self._colunas(COLUNAS_DA_ALTERACAO)
            abas.append(Aba(ALTERACOES, cabecalho, colunas, alteracoes))
        return abas

    def _colunas(self, colunas: list[Coluna]) -> list[Coluna]:
        return [*colunas, FILIAL] if self.com_filial else colunas

    def _linha(self, mov: Movimentacao, colunas: list[Coluna]) -> list:
        valores = {
            "SOLICITANTE": mov.solicitante.nome if mov.solicitante else "",
            "STATUS": STATUS_NA_PLANILHA[mov.etapa],
            "CADASTRO": cadastro(mov),
            "NOME": mov.beneficiario.nome,
            "CARTEIRINHA": mov.carteirinha,
            "DATA DIGITAÇÃO": mov.digitada_em,
            "INICIO": mov.vigencia,
            "PLANO": mov.contrato.nome_do_plano,
            "OBSERVAÇÃO": observacao(mov),
            FILIAL.titulo: str(mov.beneficiario.unidade_efetiva or ""),
        }
        return [valores[coluna.titulo] for coluna in self._colunas(colunas)]

    @property
    def nome_do_arquivo(self) -> str:
        mes = MESES[self.inicio.month - 1]
        return f"inclusoes-{slugify(mes)}-{self.inicio.year}-{slugify(self.cliente.nome)}.xlsx"

    def gerar(self) -> bytes:
        livro = Workbook()
        livro.remove(livro.active)
        abas = self.abas() or [
            # Cliente sem nenhum plano: a planilha sai com uma aba vazia.
            Aba("INCLUSÕES", f"{self.cliente} · {self.mes}", self._colunas(COLUNAS), [])
        ]
        for aba in abas:
            self._escrever_aba(livro.create_sheet(aba.titulo), aba)
        arquivo = BytesIO()
        livro.save(arquivo)
        return arquivo.getvalue()

    def _escrever_aba(self, folha, aba: Aba) -> None:
        # Linha 1: título da aba; linha 2: nomes das colunas; dados a partir da linha 3.
        folha.append([aba.cabecalho])
        folha.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(aba.colunas))
        folha["A1"].font = Font(bold=True, size=12)
        folha["A1"].alignment = Alignment(horizontal="center")
        folha.append([coluna.titulo for coluna in aba.colunas])
        for celula in folha[2]:
            celula.font = Font(bold=True)
        for linha in aba.linhas:
            folha.append(linha)
        for numero, coluna in enumerate(aba.colunas, start=1):
            folha.column_dimensions[get_column_letter(numero)].width = coluna.largura
            for (celula,) in folha.iter_rows(min_row=3, min_col=numero, max_col=numero):
                self._formatar(celula, coluna)
        folha.freeze_panes = "A3"
        ultima = get_column_letter(len(aba.colunas))
        folha.auto_filter.ref = f"A2:{ultima}{max(folha.max_row, 2)}"

    @staticmethod
    def _formatar(celula, coluna: Coluna) -> None:
        if coluna.data:
            celula.number_format = "DD/MM/YYYY"
        elif isinstance(celula.value, str) and celula.value.startswith("="):
            # Texto que começa com "=" (ex.: numa observação) fica como texto, nunca como fórmula.
            celula.data_type = "s"
