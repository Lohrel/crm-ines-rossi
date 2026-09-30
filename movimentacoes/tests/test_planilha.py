from datetime import date, datetime
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apolices.models import Contrato, Seguradora
from beneficiarios.models import Beneficiario
from beneficiarios.tests.fixtures import HOJE, contrato, incluir
from clientes.tests.fixtures import CPF_1
from movimentacoes.models import Etapa, Movimentacao
from movimentacoes.planilha import ALTERACOES, COLUNAS, FILIAL, abas_por_ramo
from movimentacoes.tests.test_movimentacoes import pedido

pytestmark = pytest.mark.django_db

URL = reverse("movimentacoes:planilha")
TITULOS = [coluna.titulo for coluna in COLUNAS]


def planilha(resposta):
    assert resposta["Content-Type"].startswith("application/vnd.openxmlformats")
    return load_workbook(BytesIO(b"".join(resposta.streaming_content)))


def linhas(aba, *titulos) -> list[tuple]:
    """Os valores das colunas pedidas, da linha 3 em diante (1: título, 2: nomes das colunas)."""
    nomes = [celula.value for celula in aba[2]]
    posicoes = [nomes.index(titulo) for titulo in titulos]
    return [
        tuple(linha[i] for i in posicoes) for linha in aba.iter_rows(min_row=3, values_only=True)
    ]


def mudar(movimentacao, **campos):
    Movimentacao.objects.filter(pk=movimentacao.pk).update(**campos)


def setembro(cliente) -> dict:
    return {"cliente": cliente.pk, "mes": "09/2026"}


def plano(empresa, seguradora, numero, produto, unidades, ramo="Saúde"):
    """Na Unimed, cada plano (PLENO, FLEX...) é um contrato com número próprio."""
    novo = contrato(empresa, seguradora, ramo, numero, unidades)
    novo.produto = produto
    novo.save()
    return novo


def test_pagina_da_planilha(logado, empresa, pessoa, saude):
    pagina = logado.get(URL).content.decode()
    assert "Baixar planilha (Excel)" in pagina
    assert f'value="{HOJE:%m/%Y}"' in pagina
    # Só aparecem os clientes com algum contrato.
    assert "Supermercados Exemplo" in pagina and "João Exemplo" not in pagina
    # O menu Movimentações leva à planilha.
    assert URL in logado.get(reverse("movimentacoes:lista")).content.decode()


def test_planilha_de_controle_do_mes(
    logado, empresa, titular, conjuge, filho, saude, dental, contato_rh
):
    titular.matricula = "4521"
    titular.save()
    incluir(titular, saude)
    # Digitada em agosto e ainda pendente: passa para a planilha de setembro.
    pedido(titular, dental, enviada_em=date(2026, 8, 20), solicitante=contato_rh)
    concluida_em_agosto = pedido(filho, dental, enviada_em=date(2026, 8, 21))
    mudar(concluida_em_agosto, etapa=Etapa.CONCLUIDA)
    retroativa = "Inclusão retroativa em caráter de exceção"
    pedido(
        conjuge,
        saude,
        enviada_em=date(2026, 9, 4),
        vigencia=date(2026, 5, 1),
        observacao=retroativa,
    )
    processada = pedido(filho, saude, enviada_em=date(2026, 9, 15))
    mudar(processada, etapa=Etapa.CONCLUIDA, carteirinha="00320000012345678")
    rejeitada = pedido(conjuge, dental, enviada_em=date(2026, 9, 22), observacao="Refazer a DPS")
    mudar(rejeitada, etapa=Etapa.REJEITADA, motivo_rejeicao="ENT")
    pedido(conjuge, dental, enviada_em=date(2026, 10, 2), vigencia=date(2026, 11, 1))  # outubro
    pedido(titular, saude, tipo="EXC", enviada_em=date(2026, 9, 10))  # exclusão não entra

    resposta = logado.get(URL, setembro(empresa))
    assert resposta["Cache-Control"] == "private, no-store"
    assert "inclusoes-setembro-2026-supermercados-exemplo.xlsx" in resposta["Content-Disposition"]
    livro = planilha(resposta)
    # Um plano por ramo: sem aba de alterações, que não teriam para onde ir.
    assert livro.sheetnames == ["DENTAL", "SAÚDE"]
    dental_, saude_ = livro["DENTAL"], livro["SAÚDE"]
    assert saude_["A1"].value == (
        "Seguradora Exemplo · Saúde · Supermercados Exemplo · setembro de 2026"
    )
    # A empresa tem filial, então a coluna FILIAL entra no fim.
    assert [celula.value for celula in saude_[2]] == [*TITULOS, FILIAL.titulo]

    assert linhas(dental_, "SOLICITANTE", "STATUS", "CADASTRO", "NOME", "DATA DIGITAÇÃO") == [
        ("Carla RH", "Pendente", 4521, "Ana Titular", datetime(2026, 8, 20)),
        (None, "Rejeitado", "Dependente", "Bruno Cônjuge", datetime(2026, 9, 22)),
    ]
    assert linhas(dental_, "OBSERVAÇÃO")[1] == (
        "Rejeitada: Prazo da entrevista médica. Refazer a DPS",
    )
    assert linhas(saude_, "STATUS", "NOME", "CARTEIRINHA", "INICIO", "OBSERVAÇÃO") == [
        ("Pendente", "Bruno Cônjuge", None, datetime(2026, 5, 1), retroativa),
        ("Processado", "Caio Filho", "00320000012345678", datetime(2026, 10, 1), None),
    ]
    # Sem o nome do produto no contrato, a coluna PLANO traz o número dele.
    assert set(linhas(saude_, "PLANO")) == {("S-1",)}
    assert set(linhas(saude_, "FILIAL")) == {("Exemplo Centro",)}
    assert saude_["F3"].number_format == saude_["G3"].number_format == "DD/MM/YYYY"


def test_plano_vigente_sem_inclusao_tem_aba_vazia(logado, empresa, matriz, seguradora, saude):
    contrato(empresa, seguradora, "Vida", "V-1", [matriz])
    cancelado = contrato(empresa, seguradora, "Dental", "D-9", [matriz])
    Contrato.objects.filter(pk=cancelado.pk).update(situacao=Contrato.Situacao.CANCELADO)
    livro = planilha(logado.get(URL, setembro(empresa)))
    assert livro.sheetnames == ["SAÚDE", "VIDA"]
    assert livro["VIDA"].max_row == 2


def test_planos_do_mesmo_ramo_ficam_na_mesma_aba(logado, empresa, matriz, seguradora, titular):
    pleno = plano(empresa, seguradora, "U-100", "PLENO", [matriz, titular.unidade])
    flex = plano(empresa, seguradora, "U-200", "FLEX", [matriz, titular.unidade])
    dental = plano(empresa, seguradora, "D-1", "", [matriz, titular.unidade], ramo="Dental")
    conjuge = Beneficiario.objects.create(
        cliente=empresa, titular=titular, parentesco="CJ", nome="Bruno Cônjuge"
    )
    pedido(titular, flex, enviada_em=date(2026, 9, 3))
    pedido(conjuge, pleno, enviada_em=date(2026, 9, 1))
    pedido(titular, dental, enviada_em=date(2026, 9, 2))

    livro = planilha(logado.get(URL, setembro(empresa)))
    # Uma aba por ramo, em ordem de digitação, e a de alterações (há dois planos de saúde).
    assert livro.sheetnames == ["DENTAL", "SAÚDE", ALTERACOES]
    assert linhas(livro["SAÚDE"], "NOME", "PLANO") == [
        ("Bruno Cônjuge", "PLENO"),
        ("Ana Titular", "FLEX"),
    ]
    assert livro["SAÚDE"]["A1"].value.startswith("Seguradora Exemplo · Saúde · ")
    assert livro[ALTERACOES].max_row == 2


def test_ramo_com_duas_seguradoras_tem_uma_aba_por_seguradora(empresa, matriz, seguradora, saude):
    outra = Seguradora.objects.create(nome="Outra Seguradora")
    pleno = plano(empresa, outra, "U-100", "PLENO", [matriz])
    flex = plano(empresa, outra, "U-200", "FLEX", [matriz])
    vida = contrato(empresa, seguradora, "Vida", "V/1", [matriz])
    assert abas_por_ramo([saude, pleno, flex, vida]) == {
        "SAÚDE SEGURADORA EXEMPLO": [saude],
        "SAÚDE OUTRA SEGURADORA": [pleno, flex],
        "VIDA": [vida],
    }


def test_aba_de_alteracoes_de_plano(logado, empresa, matriz, seguradora, titular, contato_rh):
    titular.matricula = "4521"
    titular.save()
    pleno = plano(empresa, seguradora, "U-100", "PLENO", [matriz, titular.unidade])
    flex = plano(empresa, seguradora, "U-200", "FLEX", [matriz, titular.unidade])
    incluir(titular, pleno)
    pedido(
        titular,
        flex,
        tipo="ALT",
        contrato_anterior=pleno,
        solicitante=contato_rh,
        enviada_em=date(2026, 9, 8),
        observacao="Apartamento para enfermaria",
    )

    aba = planilha(logado.get(URL, setembro(empresa)))[ALTERACOES]
    assert aba["A1"].value == "Alterações de plano · Supermercados Exemplo · setembro de 2026"
    # Como na planilha do escritório, a aba de alterações não tem carteirinha.
    assert [celula.value for celula in aba[2]] == [
        "SOLICITANTE",
        "STATUS",
        "CADASTRO",
        "NOME",
        "DATA DIGITAÇÃO",
        "INICIO",
        "PLANO",
        "OBSERVAÇÃO",
        "FILIAL",
    ]
    assert linhas(aba, "SOLICITANTE", "STATUS", "CADASTRO", "NOME", "PLANO", "OBSERVAÇÃO") == [
        (
            "Carla RH",
            "Pendente",
            4521,
            "Ana Titular",
            "FLEX",
            "PLENO para FLEX. Apartamento para enfermaria",
        ),
    ]
    assert linhas(aba, "DATA DIGITAÇÃO", "INICIO") == [
        (datetime(2026, 9, 8), datetime(2026, 10, 1))
    ]
    # A alteração não entra na aba do ramo, que é só de inclusões.
    assert planilha(logado.get(URL, setembro(empresa)))["SAÚDE"].max_row == 2


def test_pessoa_fisica_nao_tem_coluna_filial(logado, pessoa, seguradora):
    plano = contrato(pessoa, seguradora, "Saúde", "PF-1", [])
    joao = Beneficiario.objects.create(cliente=pessoa, nome="João Exemplo", cpf=CPF_1)
    pedido(joao, plano, enviada_em=date(2026, 9, 3))
    aba = planilha(logado.get(URL, setembro(pessoa)))["SAÚDE"]
    assert [celula.value for celula in aba[2]] == TITULOS
    assert linhas(aba, "CADASTRO", "NOME") == [(None, "João Exemplo")]


def test_texto_que_parece_formula_fica_como_texto(logado, empresa, titular, saude):
    formula = '=HYPERLINK("https://exemplo.invalid","abrir")'
    pedido(titular, saude, enviada_em=date(2026, 9, 3), observacao=formula)
    celula = planilha(logado.get(URL, setembro(empresa)))["SAÚDE"]["I3"]
    assert celula.value == formula and celula.data_type == "s"


def test_empresa_e_mes_sao_obrigatorios(logado, empresa, saude):
    resposta = logado.get(URL, {"cliente": "", "mes": "13/2026"})
    pagina = resposta.content.decode()
    assert resposta["Content-Type"].startswith("text/html")
    assert "Este campo é obrigatório." in pagina and "Mês inválido" in pagina


def test_planilha_exige_login(client, empresa):
    resposta = client.get(URL, setembro(empresa))
    assert resposta.status_code == 302 and reverse("login") in resposta.url
