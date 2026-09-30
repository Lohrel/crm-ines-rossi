"""Importação da planilha mensal de controle (planilhas fictícias, no modelo do escritório)."""

import re
from datetime import date, datetime
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from beneficiarios.models import Beneficiario
from beneficiarios.tests.fixtures import incluir
from movimentacoes import importacao
from movimentacoes.importacao import (
    ATUALIZAR,
    IMPORTAR,
    PROBLEMA,
    PULAR,
    UNIDADE,
    Importacao,
    PlanilhaInvalidaError,
    ler_planilha,
)
from movimentacoes.models import Etapa, Movimentacao
from movimentacoes.planilha import PlanilhaDoMes
from movimentacoes.tests.test_alteracao_de_plano import plano
from movimentacoes.tests.test_movimentacoes import pedido

pytestmark = pytest.mark.django_db

URL = reverse("movimentacoes:importar_planilha")
# As colunas de cada aba da planilha do escritório.
UNIMED = [
    "SOLICITANTE",
    "STATUS",
    "CADASTRO",
    "NOME",
    "CARTEIRINHA",
    "DATA DIGITAÇÃO",
    "INICIO",
    "PLANO",
    "OBSERVAÇÃO",
]
DENTAL = ["SOLICITANTE", "CADASTRO", "NOME", "CARTEIRINHA", "DATA DIGITAÇÃO", "INICIO"]
ALTERACOES = [coluna for coluna in UNIMED if coluna != "CARTEIRINHA"]
DIGITADA = datetime(2026, 9, 10)
OUTUBRO = datetime(2026, 10, 1)


def xlsx(*abas) -> BytesIO:
    """Título na linha 1, nomes das colunas na 2 e dados da 3 em diante, como no escritório."""
    livro = Workbook()
    livro.remove(livro.active)
    for titulo, colunas, linhas in abas:
        folha = livro.create_sheet(titulo)
        folha.append([f"{titulo} · Supermercados Exemplo · setembro de 2026"])
        folha.append(colunas)
        for linha in linhas:
            folha.append(linha)
    arquivo = BytesIO()
    livro.save(arquivo)
    arquivo.seek(0)
    return arquivo


def aba_unimed(*linhas):
    return ("UNIMED", UNIMED, linhas)


def na_unimed(nome, cadastro, status="Pendente", carteirinha="", **extra):
    return [
        extra.get("solicitante", "Carla"),
        status,
        cadastro,
        nome,
        carteirinha,
        extra.get("digitada", DIGITADA),
        OUTUBRO,
        extra.get("plano", "PLENO"),
        extra.get("obs", ""),
    ]


def importar(cliente, usuario, *abas, escolhidas=None, gravar=True):
    linhas, _ = ler_planilha(xlsx(*abas))
    importacao = Importacao(cliente, linhas, usuario, escolhidas)
    return importacao.gravar() if gravar else importacao.previa()


def situacoes(resultados) -> list[tuple[str, str]]:
    return [(resultado.situacao, resultado.mensagem) for resultado in resultados]


@pytest.fixture
def pleno(empresa, seguradora, matriz, filial):
    return plano(empresa, seguradora, "U-100", "PLENO", [matriz, filial])


@pytest.fixture
def flex(empresa, seguradora, matriz, filial):
    return plano(empresa, seguradora, "U-200", "FLEX", [matriz, filial])


@pytest.fixture
def ana(titular):
    titular.matricula = "4521"
    titular.save()
    return titular


# ------------------------------------------------------------------
# Leitura do arquivo
# ------------------------------------------------------------------


def test_le_a_planilha_no_modelo_do_escritorio():
    arquivo = xlsx(
        aba_unimed(na_unimed("Ana  Titular", 4521), [None] * 9),
        ("DENTAL", DENTAL, [["Carla", "Dependente", "Caio Filho", 123456, 46275, "01/10/2026"]]),
        (
            "ALTERAÇÕES DE PLANO",
            ALTERACOES,
            [["Carla", "Processado", 4521, "Ana Titular", DIGITADA, OUTUBRO, "FLEX", ""]],
        ),
        ("RESUMO", ["TOTAL", "VIDAS"], [[1, 2]]),
    )
    linhas, avisos = ler_planilha(arquivo)
    assert [(linha.aba, linha.numero, linha.alteracao, linha.tem_status) for linha in linhas] == [
        ("UNIMED", 3, False, True),
        ("DENTAL", 3, False, False),
        ("ALTERAÇÕES DE PLANO", 3, True, True),
    ]
    ana, caio, _ = linhas
    assert (ana["cadastro"], ana["nome"], ana["plano"]) == ("4521", "Ana Titular", "PLENO")
    assert (ana["digitada_em"], ana["vigencia"]) == ("10/09/2026", "01/10/2026")
    # Data guardada como número do Excel, e carteirinha como número.
    assert caio.dependente and caio["digitada_em"] == "10/09/2026"
    assert caio["carteirinha"] == "123456" and caio["status"] == ""
    assert avisos == ['Aba "RESUMO" ficou de fora: não tem a linha com NOME.']


def test_arquivo_que_nao_e_planilha():
    with pytest.raises(PlanilhaInvalidaError, match="Salve-o como Excel"):
        ler_planilha(BytesIO(b"NOME;CADASTRO\nAna;4521\n"))
    with pytest.raises(PlanilhaInvalidaError, match="Nenhuma linha com NOME"):
        ler_planilha(xlsx(aba_unimed()))


def test_planilha_grande_demais(monkeypatch):
    monkeypatch.setattr(importacao, "LIMITE_DESCOMPACTADO", 1000)
    with pytest.raises(PlanilhaInvalidaError, match="grande demais"):
        ler_planilha(xlsx(aba_unimed(na_unimed("Ana Titular", 4521))))
    monkeypatch.setattr(importacao, "LIMITE_DE_LINHAS", 1)
    monkeypatch.setattr(importacao, "LIMITE_DESCOMPACTADO", 10**7)
    with pytest.raises(PlanilhaInvalidaError, match="passa de 1 linhas"):
        ler_planilha(xlsx(aba_unimed(na_unimed("Ana", 1), na_unimed("Bia", 2))))


# ------------------------------------------------------------------
# O que acontece com cada linha
# ------------------------------------------------------------------


def test_previa_nao_grava_nada(empresa, usuario, ana, pleno):
    resultados = importar(
        empresa, usuario, aba_unimed(na_unimed("Ana Titular", 4521)), gravar=False
    )
    assert situacoes(resultados) == [(IMPORTAR, "Inclusão em PLENO, em análise.")]
    assert not Movimentacao.objects.exists()


def test_pendente_fica_em_analise(empresa, usuario, ana, pleno, contato_rh):
    obs = "Aguardando a entrevista médica"
    importar(empresa, usuario, aba_unimed(na_unimed("Ana Titular", 4521, obs=obs)))
    mov = Movimentacao.objects.get()
    assert (mov.tipo, mov.etapa, mov.contrato, mov.beneficiario) == ("INC", "ANA", pleno, ana)
    assert mov.recebida_em == mov.enviada_em == date(2026, 9, 10)
    assert mov.vigencia == date(2026, 10, 1) and mov.observacao == obs
    # SOLICITANTE traz só o primeiro nome da Carla, do RH.
    assert mov.solicitante == contato_rh and mov.registrada_por == usuario
    assert not ana.vinculos.exists()


def test_processado_entra_no_plano_com_a_carteirinha(
    empresa, usuario, ana, conjuge, filho, pleno, dental
):
    resultados = importar(
        empresa,
        usuario,
        # O dependente vem antes na planilha, mas o titular é importado primeiro.
        aba_unimed(
            na_unimed("Bruno Cônjuge", "Dependente", "Processado", "0032000000456"),
            na_unimed("Ana Titular", 4521, "Processado", 32000000123),
        ),
        # O DENTAL não tem STATUS: com carteirinha, a inclusão está feita.
        (
            "DENTAL",
            DENTAL,
            [
                ["Carla", 4521, "Ana Titular", 777, DIGITADA, OUTUBRO],
                ["Carla", "Dependente", "Caio Filho", None, DIGITADA, OUTUBRO],
            ],
        ),
    )
    feita = "concluída: no plano desde 01/10/2026."
    assert situacoes(resultados) == [
        (IMPORTAR, f"Inclusão em PLENO, {feita}"),
        (IMPORTAR, f"Inclusão em PLENO, {feita}"),
        (IMPORTAR, f"Inclusão em D-1, {feita}"),
        (IMPORTAR, "Inclusão em D-1, em análise."),
    ]
    assert ana.vinculos.get(contrato=pleno).carteirinha == "32000000123"
    assert conjuge.vinculos.get(contrato=pleno).data_inclusao == date(2026, 10, 1)
    mov = Movimentacao.objects.get(beneficiario=ana, contrato=pleno)
    # A planilha não diz quando a carteirinha foi enviada.
    assert mov.etapa == Etapa.CONCLUIDA and mov.carteirinha_enviada_em is None
    assert not filho.vinculos.exists()


def test_lembrete_na_coluna_da_carteirinha(empresa, usuario, ana, pleno):
    linha = na_unimed("Ana Titular", 4521, "Processado", "INCLUIR NO DENTAL")
    resultados = importar(empresa, usuario, aba_unimed(linha))
    assert situacoes(resultados) == [
        (IMPORTAR, "Inclusão em PLENO, aprovada: falta a carteirinha para concluir.")
    ]
    mov = Movimentacao.objects.get()
    assert (mov.etapa, mov.carteirinha) == (Etapa.APROVADA, "")
    assert mov.observacao == "Lembrete da planilha: INCLUIR NO DENTAL."
    assert ana.vinculos.get().data_inclusao == date(2026, 10, 1)


def test_titular_achado_pelo_nome_ou_pela_matricula(empresa, usuario, titular, pleno, flex):
    # Sem matrícula no sistema: achada pelo nome, e a matrícula é preenchida.
    [resultado] = importar(
        empresa, usuario, aba_unimed(na_unimed("ANA TITULAR", 4521, solicitante="Zeca"))
    )
    assert resultado.situacao == IMPORTAR
    assert resultado.avisos == [
        "Solicitante Zeca: nenhum contato com este nome; fica em branco.",
        "Matrícula preenchida no cadastro da pessoa.",
    ]
    titular.refresh_from_db()
    assert titular.matricula == "4521"
    # O Excel tira o zero da frente: 04521 e 4521 são a mesma matrícula.
    titular.matricula = "04521"
    titular.save()
    linha = na_unimed("Ana T.", 4521, plano="FLEX", solicitante="")
    [resultado] = importar(empresa, usuario, aba_unimed(linha))
    assert resultado.situacao == IMPORTAR
    assert resultado.avisos == ["No sistema, a matrícula 4521 é de Ana Titular."]
    assert Movimentacao.objects.filter(beneficiario=titular).count() == 2


def test_titular_novo_entra_no_cadastro(empresa, usuario, matriz, filial, pleno, contato_rh):
    [resultado] = importar(empresa, usuario, aba_unimed(na_unimed("Dora Nova", 9001)))
    assert resultado.situacao == IMPORTAR
    # A Carla, do RH, atende só a filial: a pessoa entra nela, e a prévia deixa trocar.
    assert resultado.avisos == [
        "Titular novo: cadastrado só com nome e matrícula, na unidade Exemplo Centro."
    ]
    assert resultado.unidades == [matriz, filial] and resultado.escolhida == filial.pk
    dora = Beneficiario.objects.get(nome="Dora Nova")
    assert (dora.matricula, dora.unidade, dora.eh_titular) == ("9001", filial, True)
    assert Movimentacao.objects.get().beneficiario == dora


def test_titular_novo_sem_unidade_conhecida(empresa, usuario, matriz, filial, pleno):
    aba = aba_unimed(na_unimed("Dora Nova", 9001, solicitante=""))
    [resultado] = importar(empresa, usuario, aba, gravar=False)
    assert (resultado.situacao, resultado.escolhida) == (UNIDADE, None)
    assert resultado.mensagem == "Titular novo: escolha a unidade em que trabalha."
    escolhidas = {resultado.linha.id: matriz.pk}
    [resultado] = importar(empresa, usuario, aba, escolhidas=escolhidas)
    assert resultado.situacao == IMPORTAR and resultado.escolhida == matriz.pk
    assert Beneficiario.objects.get(nome="Dora Nova").unidade == matriz


def test_titular_novo_em_duas_abas_escolhe_a_unidade_uma_vez(
    empresa, usuario, matriz, filial, pleno, dental
):
    abas = [
        aba_unimed(na_unimed("Dora Nova", 9001, solicitante="")),
        ("DENTAL", DENTAL, [["", 9001, "Dora Nova", 555, DIGITADA, OUTUBRO]]),
    ]
    unimed, odonto = importar(empresa, usuario, *abas, gravar=False)
    assert unimed.situacao == UNIDADE and unimed.unidades == [matriz, filial]
    assert odonto.situacao == UNIDADE and odonto.unidades == []
    assert odonto.mensagem == "Titular novo: escolha a unidade na linha 3 da aba UNIMED."
    resultados = importar(empresa, usuario, *abas, escolhidas={unimed.linha.id: filial.pk})
    assert [resultado.situacao for resultado in resultados] == [IMPORTAR, IMPORTAR]
    dora = Beneficiario.objects.get(nome="Dora Nova")
    assert dora.unidade == filial and dora.vinculos.get().contrato == dental


def test_linhas_com_problema_ficam_de_fora(empresa, usuario, ana, pleno):
    resultados = importar(
        empresa,
        usuario,
        aba_unimed(
            na_unimed("Eva Sem Cadastro", "Dependente"),
            na_unimed("Ana Titular", 4521, plano="MASTER"),
            na_unimed("Ana Titular", 4521, digitada="31/02/2026"),
            na_unimed("Ana Titular", 4521, digitada=None),
            na_unimed("Ana Titular", 4521, status="Rejeitado"),
            na_unimed("Dora Nova", 9001, plano="", solicitante=""),
        ),
    )
    assert situacoes(resultados) == [
        (PROBLEMA, "Dependente sem cadastro: cadastre pela página do titular e importe de novo."),
        (PROBLEMA, "Plano MASTER não encontrado: confira o Produto/plano dos contratos."),
        (PROBLEMA, "DATA DIGITAÇÃO inválida: 31/02/2026."),
        (PROBLEMA, "Falta a DATA DIGITAÇÃO."),
        (PULAR, "STATUS Rejeitado: só Pendente e Processado são importados."),
        (PROBLEMA, 'Informe o PLANO: a aba "UNIMED" tem mais de um plano.'),
    ]
    assert not Movimentacao.objects.exists()
    assert not Beneficiario.objects.filter(nome="Dora Nova").exists()


def test_plano_de_contrato_cancelado(empresa, usuario, ana, pleno):
    pleno.situacao = "CAN"
    pleno.save()
    resultados = importar(empresa, usuario, aba_unimed(na_unimed("Ana Titular", 4521)))
    assert situacoes(resultados) == [(PROBLEMA, "O contrato do plano PLENO está cancelado.")]


def test_importar_de_novo_nao_repete(empresa, usuario, ana, pleno, contato_rh):
    aba = aba_unimed(
        na_unimed("Ana Titular", 4521, "Processado", "0032000000123"),
        na_unimed("Dora Nova", 9001),
    )
    importar(empresa, usuario, aba)
    resultados = importar(empresa, usuario, aba)
    assert situacoes(resultados) == [
        (PULAR, "Já está no plano PLENO."),
        (PULAR, "Já registrada, em análise."),
    ]
    assert Movimentacao.objects.count() == 2
    assert Beneficiario.objects.filter(nome="Dora Nova").count() == 1


def test_pedido_em_aberto_que_a_planilha_da_como_processado_avanca(empresa, usuario, ana, pleno):
    importar(empresa, usuario, aba_unimed(na_unimed("Ana Titular", 4521)))
    processado = na_unimed("Ana Titular", 4521, "Processado", "0032000000123")
    resultados = importar(empresa, usuario, aba_unimed(processado))
    assert situacoes(resultados) == [
        (ATUALIZAR, "Já registrada, em análise; agora concluída: no plano desde 01/10/2026.")
    ]
    mov = Movimentacao.objects.get()
    assert (mov.etapa, mov.carteirinha) == (Etapa.CONCLUIDA, "0032000000123")
    assert ana.vinculos.get().carteirinha == "0032000000123"


def test_alteracao_de_plano(
    empresa, usuario, seguradora, matriz, filial, ana, conjuge, pleno, flex
):
    amigo = plano(empresa, seguradora, "U-300", "AMIGO", [matriz, filial])
    incluir(ana, pleno, inclusao=date(2026, 1, 1))
    incluir(ana, amigo, inclusao=date(2026, 1, 1))
    incluir(conjuge, pleno, inclusao=date(2026, 1, 1))
    linha = ["Carla", "Processado", 4521, "Ana Titular", DIGITADA, OUTUBRO, "FLEX"]
    # Com dois planos de saúde, a observação diz de qual a pessoa sai.
    aba = ("ALTERAÇÕES DE PLANO", ALTERACOES, [[*linha, "Pleno para Flex Apartamento"]])
    resultados = importar(empresa, usuario, aba)
    assert situacoes(resultados) == [
        (IMPORTAR, "Alteração de PLENO para FLEX, concluída: no plano desde 01/10/2026.")
    ]
    mov = Movimentacao.objects.get()
    assert (mov.tipo, mov.contrato_anterior, mov.contrato) == ("ALT", pleno, flex)
    # O cônjuge muda junto, e o AMIGO continua.
    assert conjuge.vinculos.get(data_exclusao__isnull=True).contrato == flex
    assert ana.vinculos.get(contrato=amigo).data_exclusao is None


def test_inclusao_de_quem_ja_tem_plano_do_ramo(empresa, usuario, ana, pleno, flex, contato_rh):
    incluir(ana, pleno)
    [resultado] = importar(
        empresa, usuario, aba_unimed(na_unimed("Ana Titular", 4521, plano="FLEX"))
    )
    assert resultado.situacao == IMPORTAR
    assert resultado.avisos == [
        "A pessoa já está no plano PLENO, de Saúde. Se é troca de plano, lance na aba "
        "ALTERAÇÕES DE PLANO."
    ]


def test_carteirinha_comprida_demais_so_afeta_a_linha(empresa, usuario, ana, conjuge, pleno):
    pedido(ana, pleno, enviada_em=date(2026, 9, 10))
    comprida = "9" * 61
    resultados = importar(
        empresa,
        usuario,
        aba_unimed(
            na_unimed("Ana Titular", 4521, "Processado", comprida),
            na_unimed("Bruno Cônjuge", "Dependente"),
        ),
    )
    assert [resultado.situacao for resultado in resultados] == [PROBLEMA, IMPORTAR]
    assert "no máximo 60 caracteres" in resultados[0].mensagem
    assert Movimentacao.objects.get(beneficiario=ana).etapa == Etapa.EM_ANALISE


def test_planilha_baixada_do_sistema_volta_e_faz_o_pedido_avancar(
    empresa, usuario, ana, pleno, flex, contato_rh
):
    pedido(ana, pleno, enviada_em=date(2026, 9, 10), solicitante=contato_rh)
    baixada = BytesIO(PlanilhaDoMes(empresa, date(2026, 9, 1)).gerar())
    resultados = Importacao(empresa, ler_planilha(baixada)[0], usuario).gravar()
    assert situacoes(resultados) == [(PULAR, "Já registrada, em análise.")]
    # Na planilha, o escritório marca o pedido como processado e anota a carteirinha.
    livro = load_workbook(baixada)
    folha = livro["SAÚDE"]
    colunas = [celula.value for celula in folha[2]]
    folha.cell(3, colunas.index("STATUS") + 1, "Processado")
    folha.cell(3, colunas.index("CARTEIRINHA") + 1, "0032000000123")
    arquivo = BytesIO()
    livro.save(arquivo)
    resultados = Importacao(empresa, ler_planilha(arquivo)[0], usuario).gravar()
    assert [resultado.situacao for resultado in resultados] == [ATUALIZAR]
    assert Movimentacao.objects.get().etapa == Etapa.CONCLUIDA


# ------------------------------------------------------------------
# Telas
# ------------------------------------------------------------------


def enviar(logado, cliente, arquivo, nome="inclusoes-setembro.xlsx") -> str:
    conteudo = arquivo.getvalue() if isinstance(arquivo, BytesIO) else arquivo
    upload = SimpleUploadedFile(nome, conteudo)
    return logado.post(URL, {"cliente": cliente.pk, "arquivo": upload}).content.decode()


def dados_da_previa(pagina: str) -> str:
    return re.search(r'name="dados" value="([^"]+)"', pagina).group(1)


def test_tela_de_importacao(logado, empresa, pleno):
    pagina = logado.get(URL).content.decode()
    assert "Ver prévia" in pagina and 'accept=".xlsx"' in pagina
    assert "Supermercados Exemplo" in pagina
    # A página da planilha mensal leva à importação.
    assert URL in logado.get(reverse("movimentacoes:planilha")).content.decode()


def test_enviar_ver_a_previa_e_confirmar(logado, empresa, ana, pleno, contato_rh):
    pagina = enviar(logado, empresa, xlsx(aba_unimed(na_unimed("Ana Titular", 4521))))
    assert "Prévia da importação" in pagina and "Nada foi gravado ainda" in pagina
    assert "Inclusão em PLENO, em análise." in pagina and "1 a importar" in pagina
    assert "Confirmar importação" in pagina and not Movimentacao.objects.exists()

    pagina = logado.post(URL, {"dados": dados_da_previa(pagina), "acao": "confirmar"})
    pagina = pagina.content.decode()
    assert "Planilha importada" in pagina and "1 importada" in pagina
    assert f"?cliente={empresa.pk}&amp;situacao=todas" in pagina
    assert Movimentacao.objects.get().registrada_por.username == "ana"


def test_confirmar_pede_a_unidade_do_titular_novo(logado, empresa, matriz, filial, pleno):
    linha = na_unimed("Dora Nova", 9001, solicitante="")
    pagina = enviar(logado, empresa, xlsx(aba_unimed(linha)))
    assert '<select name="unidade-0-3"' in pagina and "Atualizar prévia" in pagina
    dados = dados_da_previa(pagina)

    pagina = logado.post(URL, {"dados": dados, "acao": "confirmar"}).content.decode()
    assert "Escolha a unidade de cada titular novo antes de confirmar." in pagina
    assert not Beneficiario.objects.filter(nome="Dora Nova").exists()

    escolha = {"dados": dados, "unidade-0-3": matriz.pk}
    pagina = logado.post(URL, {**escolha, "acao": "previa"}).content.decode()
    assert f'<option value="{matriz.pk}" selected>' in pagina
    pagina = logado.post(URL, {**escolha, "acao": "confirmar"}).content.decode()
    assert "Planilha importada" in pagina
    assert Beneficiario.objects.get(nome="Dora Nova").unidade == matriz


def test_nada_a_importar(logado, empresa, ana, pleno):
    incluir(ana, pleno)
    pagina = enviar(logado, empresa, xlsx(aba_unimed(na_unimed("Ana Titular", 4521))))
    assert "Nenhuma linha desta planilha muda algo no sistema." in pagina
    assert "Confirmar importação" not in pagina and "Enviar outra planilha" in pagina


def test_previa_adulterada_ou_vencida(logado):
    resposta = logado.post(URL, {"dados": "adulterado", "acao": "confirmar"}, follow=True)
    assert resposta.redirect_chain == [(URL, 302)]
    assert "A prévia expirou. Envie a planilha de novo." in resposta.content.decode()


def test_arquivo_invalido_na_tela(logado, empresa, pleno):
    pagina = enviar(logado, empresa, b"isto nao e uma planilha")
    assert "Não foi possível ler o arquivo" in pagina
    pagina = enviar(logado, empresa, b"NOME;CADASTRO", nome="inclusoes.csv")
    assert "As extensões válidas são: xlsx" in pagina


def test_importacao_exige_login(client):
    assert client.get(URL).status_code == 302
