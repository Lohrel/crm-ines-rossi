"""Importação da planilha mensal de controle, no formato que o escritório já usa.

Cada linha das abas de inclusão (ex.: UNIMED, SAÚDE, DENTAL) vira um pedido de
inclusão, e cada linha da aba ALTERAÇÕES DE PLANO, um pedido de alteração. A
pessoa é achada pela matrícula (titular) ou pelo nome (dependente); o plano,
pela coluna PLANO (um plano cadastrado no contrato, como TNPQ, ou o "Produto/plano"
do contrato); o solicitante, pelo primeiro nome entre os contatos da empresa.

A prévia faz tudo dentro de uma transação desfeita no fim, para mostrar
exatamente o que vai acontecer; ao confirmar, faz de novo, agora gravando.
Importar a mesma planilha outra vez não repete nada: o que já está no sistema
fica como está, e o pedido que estava em aberto e na planilha já aparece como
Processado avança.
"""

import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta

from django.core import signing
from django.core.exceptions import ValidationError
from django.db import DatabaseError, transaction
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apolices.models import Contrato, Plano
from beneficiarios.models import Beneficiario, Vinculo
from clientes.models import Cliente, Contato, Unidade
from clientes.planilhas import (
    NAO_LEU,
    PlanilhaInvalidaError,
    chave,
    conferir_tamanho,
    mesmo_codigo,
)
from movimentacoes.models import Etapa, Movimentacao

# Nome de cada coluna da planilha (sem acento, em maiúsculas) e o dado que ela traz.
COLUNAS = {
    "SOLICITANTE": "solicitante",
    "STATUS": "status",
    "ETAPA": "status",
    "CADASTRO": "cadastro",
    "MATRICULA": "cadastro",
    "NOME": "nome",
    "CARTEIRINHA": "carteirinha",
    "DATA DIGITACAO": "digitada_em",
    "DATA DE DIGITACAO": "digitada_em",
    "DATA DE ENVIO": "digitada_em",
    "INICIO": "vigencia",
    "INICIO DE VIGENCIA": "vigencia",
    "VIGENCIA": "vigencia",
    "PLANO": "plano",
    "OBSERVACAO": "observacao",
    "OBS": "observacao",
    "FILIAL": "filial",
}
# Os dados que as colunas trazem.
DADOS = set(COLUNAS.values())
DATAS = {"digitada_em", "vigencia"}
FEITAS = {"PROCESSADO", "PROCESSADA", "CONCLUIDO", "CONCLUIDA"}
PENDENTES = {"PENDENTE", "EM ANALISE"}
# Aba das alterações de plano (ex.: ALTERAÇÕES DE PLANO); as outras são de inclusão.
ABAS_DE_ALTERACAO = ("ALTERA", "MUDANCA", "TROCA")
LIMITE_DE_LINHAS = 2000
# O .xlsx é um .zip: descompactado, uma planilha mensal fica bem abaixo disto.
LIMITE_DESCOMPACTADO = 20 * 1024 * 1024
VALIDADE_DA_PREVIA = timedelta(hours=2)

# O que acontece com cada linha.
IMPORTAR = "importar"
ATUALIZAR = "atualizar"
PULAR = "pular"
UNIDADE = "unidade"
PROBLEMA = "problema"


def _rotulo(contrato: Contrato) -> str:
    """Como o plano pode ser escrito na coluna PLANO: o produto ou o número do contrato."""
    return f"{contrato.produto} ou {contrato.numero}" if contrato.produto else contrato.numero


@dataclass
class Linha:
    aba: str
    indice_da_aba: int
    numero: int
    alteracao: bool
    tem_status: bool
    valores: dict[str, str]

    @property
    def id(self) -> str:
        return f"{self.indice_da_aba}-{self.numero}"

    def __getitem__(self, nome: str) -> str:
        # Só os dados das colunas; o resto (ex.: numero, no template) continua atributo.
        if nome not in DADOS:
            raise KeyError(nome)
        return self.valores.get(nome, "")

    @property
    def dependente(self) -> bool:
        return chave(self["cadastro"]) in ("DEPENDENTE", "DEP")


def _texto(valor, coluna: str) -> str:
    if valor is None:
        return ""
    if coluna in DATAS:
        if isinstance(valor, (int, float)) and 20000 < valor < 80000:
            valor = from_excel(valor)  # data guardada como número
        if isinstance(valor, (datetime, date)):
            return f"{valor:%d/%m/%Y}"
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _cabecalho(linhas: list[tuple]) -> tuple[int, dict[int, str]] | None:
    """A linha dos títulos (NOME, CADASTRO...) entre as primeiras da aba."""
    for numero, valores in enumerate(linhas, start=1):
        colunas = {i: COLUNAS[chave(v)] for i, v in enumerate(valores) if chave(v) in COLUNAS}
        if "nome" in colunas.values() and len(colunas) >= 3:
            return numero, colunas
    return None


def ler_planilha(arquivo) -> tuple[list[Linha], list[str]]:
    """As linhas com nome de todas as abas, e avisos sobre as abas que ficaram de fora."""
    conferir_tamanho(arquivo, LIMITE_DESCOMPACTADO)
    try:
        livro = load_workbook(arquivo, read_only=True, data_only=True, keep_links=False)
        try:
            linhas, avisos = _ler_abas(livro)
        finally:
            livro.close()
    except PlanilhaInvalidaError:
        raise
    except Exception as erro:  # o openpyxl levanta erros de vários tipos com arquivo estragado
        raise PlanilhaInvalidaError(NAO_LEU) from erro
    if not linhas:
        raise PlanilhaInvalidaError("Nenhuma linha com NOME foi encontrada na planilha.")
    return linhas, avisos


def _ler_abas(livro) -> tuple[list[Linha], list[str]]:
    linhas, avisos = [], []
    for indice, folha in enumerate(livro.worksheets):
        cabecalho = _cabecalho(list(folha.iter_rows(max_row=10, values_only=True)))
        if cabecalho is None:
            avisos.append(f'Aba "{folha.title}" ficou de fora: não tem a linha com NOME.')
            continue
        numero_do_cabecalho, colunas = cabecalho
        tem_status = "status" in colunas.values()
        alteracao = any(palavra in chave(folha.title) for palavra in ABAS_DE_ALTERACAO)
        dados = folha.iter_rows(min_row=numero_do_cabecalho + 1, values_only=True)
        for numero, valores in enumerate(dados, start=numero_do_cabecalho + 1):
            lidos = {
                nome: _texto(valores[i], nome) for i, nome in colunas.items() if i < len(valores)
            }
            if not lidos.get("nome"):
                continue
            linhas.append(Linha(folha.title, indice, numero, alteracao, tem_status, lidos))
            if len(linhas) > LIMITE_DE_LINHAS:
                raise PlanilhaInvalidaError(f"A planilha passa de {LIMITE_DE_LINHAS} linhas.")
    return linhas, avisos


def empacotar(cliente: Cliente, arquivo: str, linhas: list[Linha], avisos: list[str]) -> str:
    """As linhas lidas, assinadas, para a confirmação não depender de reenviar o arquivo."""
    dados = {
        "cliente": cliente.pk,
        "arquivo": arquivo,
        "linhas": [asdict(linha) for linha in linhas],
        "avisos": avisos,
    }
    return signing.dumps(dados, salt="movimentacoes.importacao", compress=True)


def desempacotar(texto: str) -> dict:
    """Levanta signing.BadSignature se o conteúdo foi alterado ou a prévia expirou."""
    dados = signing.loads(
        texto, salt="movimentacoes.importacao", max_age=VALIDADE_DA_PREVIA.total_seconds()
    )
    dados["linhas"] = [Linha(**linha) for linha in dados["linhas"]]
    return dados


@dataclass
class Resultado:
    linha: Linha
    situacao: str
    mensagem: str
    avisos: list[str] = field(default_factory=list)
    # Titular novo numa empresa com mais de uma unidade: as opções e a escolhida.
    unidades: list[Unidade] = field(default_factory=list)
    escolhida: int | None = None


class NaoImportarError(Exception):
    """A linha já está no sistema (ou não é importada): fica como está."""


class FaltaUnidadeError(Exception):
    """Titular novo sem como saber a unidade: a escolha fica para a prévia."""


class Importacao:
    """Aplica as linhas lidas da planilha no cadastro de um cliente."""

    def __init__(self, cliente: Cliente, linhas: list[Linha], usuario, unidades_escolhidas=None):
        self.cliente = cliente
        self.linhas = linhas
        self.usuario = usuario
        self.unidades = list(cliente.unidades.all()) if cliente.eh_empresa else []
        por_pk = {unidade.pk: unidade for unidade in self.unidades}
        self.unidades_escolhidas = {
            linha: por_pk[pk] for linha, pk in (unidades_escolhidas or {}).items() if pk in por_pk
        }
        self.contratos = list(
            Contrato.objects.de_pessoas()
            .filter(cliente=cliente)
            .select_related("ramo", "seguradora")
            .prefetch_related("planos", "unidades")
        )
        self.contatos = (
            list(Contato.objects.filter(cliente=cliente).prefetch_related("unidades"))
            if cliente.eh_empresa
            else []
        )

    def previa(self) -> list[Resultado]:
        """O que aconteceria com cada linha, sem gravar nada."""
        with transaction.atomic():
            resultados = self._processar()
            transaction.set_rollback(True)
        return resultados

    def gravar(self) -> list[Resultado]:
        with transaction.atomic():
            # Duas importações do mesmo cliente ao mesmo tempo (ex.: clique duplo) vão em fila.
            Cliente.objects.select_for_update().get(pk=self.cliente.pk)
            return self._processar()

    def _processar(self) -> list[Resultado]:
        self._carregar_pessoas()
        # Unidade de cada titular novo, escolhida na primeira linha em que ele aparece.
        self.unidades_novas: dict[str, tuple[Linha, Unidade | None]] = {}
        # Titulares antes dos dependentes, e inclusões antes das alterações.
        ordem = sorted(self.linhas, key=lambda linha: (linha.alteracao, linha.dependente))
        resultados = {linha.id: self._processar_linha(linha) for linha in ordem}
        return [resultados[linha.id] for linha in self.linhas]

    def _carregar_pessoas(self) -> None:
        self.pessoas = list(
            Beneficiario.objects.filter(cliente=self.cliente).select_related("titular")
        )

    def _processar_linha(self, linha: Linha) -> Resultado:
        self.avisos, self.opcoes, self.escolhida = [], [], None
        try:
            with transaction.atomic():
                situacao, mensagem = self._importar(linha)
        except (NaoImportarError, FaltaUnidadeError, ValidationError, DatabaseError) as motivo:
            # A linha foi desfeita no banco: as pessoas em memória voltam a ser as do banco.
            self._carregar_pessoas()
            situacao, mensagem = self._sem_importar(motivo)
        return Resultado(linha, situacao, mensagem, self.avisos, self.opcoes, self.escolhida)

    @staticmethod
    def _sem_importar(motivo: Exception) -> tuple[str, str]:
        if isinstance(motivo, NaoImportarError):
            return PULAR, str(motivo)
        if isinstance(motivo, FaltaUnidadeError):
            return UNIDADE, str(motivo)
        if isinstance(motivo, ValidationError):
            return PROBLEMA, " ".join(motivo.messages)
        return PROBLEMA, "Não foi possível gravar esta linha. Confira a pessoa e o plano."

    # ------------------------------------------------------------------
    # Uma linha
    # ------------------------------------------------------------------

    def _importar(self, linha: Linha) -> tuple[str, str]:
        etapa = self._etapa(linha)
        digitada_em = self._data(linha, "digitada_em", "DATA DIGITAÇÃO")
        if digitada_em is None:
            raise ValidationError("Falta a DATA DIGITAÇÃO.")
        vigencia = self._data(linha, "vigencia", "INICIO")
        candidatos, ambiguo = self._plano(linha)
        solicitante = self._solicitante(linha)
        pessoa = self._pessoa(linha, solicitante)
        contrato, plano = self._escolher_pela_pessoa(candidatos, pessoa, ambiguo)
        falta_o_plano = ""
        if plano is None:
            plano, falta_o_plano = self._unico_plano_do_contrato(contrato)
        carteirinha, observacao = self._carteirinha_e_observacao(linha)
        if plano and not linha.alteracao and (anotado := self._anotar_plano(pessoa, plano)):
            return anotado
        pedido = Movimentacao(
            tipo=Movimentacao.Tipo.ALTERACAO if linha.alteracao else Movimentacao.Tipo.INCLUSAO,
            cliente=self.cliente,
            beneficiario=pessoa,
            contrato=contrato,
            plano=plano,
            solicitante=solicitante,
            recebida_em=digitada_em,
            enviada_em=digitada_em,
            vigencia=vigencia or contrato.vigencia_sugerida(digitada_em),
            carteirinha="" if linha.alteracao else carteirinha,
            observacao=observacao,
            registrada_por=self.usuario,
        )
        registrado = self._ja_registrado(pedido)
        if registrado:
            return self._atualizar(registrado, etapa, pedido.carteirinha)
        if falta_o_plano:
            # Só agora: quem já está no contrato ou já tem o pedido não precisa do plano.
            raise ValidationError(falta_o_plano)
        if not contrato.vigente:
            situacao = contrato.get_situacao_display().lower()
            raise ValidationError(f"O contrato do plano {pedido.nome_do_plano} está {situacao}.")
        if linha.alteracao:
            atual = self._plano_anterior(linha, pessoa, contrato, plano)
            pedido.contrato_anterior, pedido.plano_anterior = atual.contrato, atual.plano
        else:
            self._avisar_se_ja_tem_plano_do_ramo(pessoa, contrato)
        pedido.full_clean()
        pedido.save()
        if pedido.eh_alteracao:
            descricao = f"Alteração de {pedido.nome_do_plano_anterior} para {pedido.nome_do_plano}"
        else:
            descricao = f"Inclusão em {pedido.nome_do_plano}"
        if etapa == Etapa.EM_ANALISE:
            return IMPORTAR, f"{descricao}, em análise."
        return IMPORTAR, f"{descricao}, {self._concluir(pedido)}."

    def _atualizar(self, pedido: Movimentacao, etapa: str, carteirinha: str) -> tuple[str, str]:
        """Pedido em aberto no sistema que a planilha já dá como processado: avança."""
        antes = f"Já registrada, {pedido.get_etapa_display().lower()}"
        pedido.carteirinha = carteirinha or pedido.carteirinha
        sem_carteirinha = pedido.eh_inclusao and not pedido.carteirinha
        if etapa == Etapa.EM_ANALISE or (pedido.etapa == Etapa.APROVADA and sem_carteirinha):
            raise NaoImportarError(f"{antes}.")
        pedido.full_clean()
        return ATUALIZAR, f"{antes}; agora {self._concluir(pedido)}."

    @staticmethod
    def _concluir(pedido: Movimentacao) -> str:
        """Aprova (a pessoa entra no plano na vigência) e, com a carteirinha, conclui."""
        if pedido.etapa == Etapa.EM_ANALISE:
            pedido.mudar_etapa(Etapa.APROVADA)
        if pedido.eh_inclusao and not pedido.carteirinha:
            return "aprovada: falta a carteirinha para concluir"
        pedido.mudar_etapa(Etapa.CONCLUIDA)
        # A planilha não diz quando a carteirinha foi enviada.
        Movimentacao.objects.filter(pk=pedido.pk).update(carteirinha_enviada_em=None)
        return f"concluída: no plano desde {pedido.vigencia:%d/%m/%Y}"

    def _etapa(self, linha: Linha) -> str:
        status = chave(linha["status"])
        if status in FEITAS:
            return Etapa.CONCLUIDA
        if status in PENDENTES or (linha.tem_status and not status):
            return Etapa.EM_ANALISE
        if not linha.tem_status:
            # Aba sem STATUS (ex.: DENTAL, de inclusão imediata): com carteirinha, está feita.
            return Etapa.CONCLUIDA if self._carteirinha_e_observacao(linha)[0] else Etapa.EM_ANALISE
        raise NaoImportarError(
            f"STATUS {linha['status']}: só Pendente e Processado são importados."
        )

    @staticmethod
    def _data(linha: Linha, nome: str, coluna: str) -> date | None:
        texto = linha[nome]
        if not texto:
            return None
        for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(texto, formato).date()
            except ValueError:
                continue
        raise ValidationError(f"{coluna} inválida: {texto}.")

    @staticmethod
    def _carteirinha_e_observacao(linha: Linha) -> tuple[str, str]:
        carteirinha, observacao = linha["carteirinha"], linha["observacao"]
        if re.search(r"[A-Za-z]", carteirinha):
            # Ex.: "INCLUIR NO DENTAL" é um lembrete, não o número da carteirinha.
            lembrete = f"Lembrete da planilha: {carteirinha}."
            return "", f"{observacao} {lembrete}".strip()
        return carteirinha, observacao

    def _plano(self, linha: Linha) -> tuple[list[tuple[Contrato, Plano | None]], str]:
        """Os planos possíveis da linha, pelo PLANO (um plano do contrato, o produto ou o número
        do contrato) ou, sem ele, pela aba; e o aviso para quando sobrar mais de um."""
        nome = chave(linha["plano"])
        if nome:
            pelo_plano = self._pelo_nome_do_plano(linha, nome)
            if pelo_plano:
                return pelo_plano, (
                    f"Há mais de um plano {linha['plano']} neste cliente, em contratos diferentes "
                    f"({', '.join(c.numero for c, _ in pelo_plano)})."
                )
            achados = [c for c in self.contratos if nome in (chave(c.produto), chave(c.numero))]
            if len(achados) > 1:
                achados = self._da_aba(linha, achados) or achados
            if achados:
                return [(c, None) for c in achados], (
                    f"Há mais de um plano {linha['plano']} neste cliente: informe o número do "
                    f"contrato ({', '.join(c.numero for c in achados)})."
                )
            raise ValidationError(
                f"Plano {linha['plano']} não encontrado: confira o Produto/plano dos contratos "
                "e os planos cadastrados neles."
            )
        da_aba = self._da_aba(linha, self.contratos)
        if da_aba:
            planos = ", ".join(_rotulo(c) for c in da_aba)
            return [(c, None) for c in da_aba], (
                f'Informe o PLANO: a aba "{linha.aba}" tem mais de um plano deste cliente '
                f"({planos})."
            )
        raise ValidationError(
            f"Informe o PLANO ou cadastre o contrato: nenhum plano deste cliente é do ramo "
            f'ou da seguradora da aba "{linha.aba}".'
        )

    def _pelo_nome_do_plano(self, linha: Linha, nome: str) -> list[tuple[Contrato, Plano]]:
        """Os planos cadastrados nos contratos com este nome (ex.: TNPQ)."""
        achados = [(c, p) for c in self.contratos for p in c.planos.all() if chave(p.nome) == nome]
        if len(achados) > 1:
            da_aba = self._da_aba(linha, [c for c, _ in achados])
            achados = [(c, p) for c, p in achados if c in da_aba] or achados
        return achados

    @staticmethod
    def _escolher_pela_pessoa(
        candidatos: list[tuple[Contrato, Plano | None]], pessoa: Beneficiario, ambiguo: str
    ) -> tuple[Contrato, Plano | None]:
        """Com o mesmo plano em mais de um contrato (ex.: uma apólice por empresa do grupo),
        fica, nesta ordem: o plano em que a pessoa já está; o contrato em que ela já está (ou
        tem pedido em aberto); o que cobre a filial dela; o contrato vigente e o plano oferecido.
        """
        if len(candidatos) > 1 and pessoa.pk:
            no_plano = set(
                pessoa.vinculos.filter(data_exclusao__isnull=True).values_list("contrato", "plano")
            )
            no_contrato = {contrato for contrato, _ in no_plano} | set(
                pessoa.movimentacoes.em_aberto().values_list("contrato", flat=True)
            )
            candidatos = (
                [(c, p) for c, p in candidatos if p and (c.pk, p.pk) in no_plano]
                or [(c, p) for c, p in candidatos if c.pk in no_contrato]
                or candidatos
            )
        unidade = pessoa.unidade_efetiva
        if len(candidatos) > 1 and unidade:
            cobrem = [(c, p) for c, p in candidatos if unidade in c.unidades.all()]
            candidatos = cobrem or candidatos
        if len(candidatos) > 1:
            em_uso = [(c, p) for c, p in candidatos if c.vigente and (p is None or p.ativo)]
            candidatos = em_uso or candidatos
        if len(candidatos) > 1:
            raise ValidationError(ambiguo)
        return candidatos[0]

    @staticmethod
    def _unico_plano_do_contrato(contrato: Contrato) -> tuple[Plano | None, str]:
        """Sem o plano na linha: o único plano do contrato ou, se há mais de um, o aviso."""
        planos = contrato.planos_ativos
        if len(planos) > 1:
            nomes = ", ".join(plano.nome for plano in planos)
            return None, (
                f"O contrato {contrato.numero} tem mais de um plano: informe na coluna PLANO "
                f"qual deles ({nomes})."
            )
        return (planos[0] if planos else None), ""

    def _anotar_plano(self, pessoa: Beneficiario, plano: Plano) -> tuple[str, str] | None:
        """Quem já está no contrato sem o plano informado fica com o plano da planilha."""
        sem_plano = pessoa.vinculos.filter(
            contrato=plano.contrato_id, plano__isnull=True, data_exclusao__isnull=True
        ).first()
        if sem_plano is None:
            return None
        sem_plano.plano = plano
        sem_plano.save(update_fields=["plano", "atualizado_em"])
        return ATUALIZAR, f"Já estava no contrato; agora com o plano {plano} anotado."

    @staticmethod
    def _da_aba(linha: Linha, contratos: list[Contrato]) -> list[Contrato]:
        """Os contratos cujo ramo ou seguradora aparece no nome da aba (ex.: DENTAL, UNIMED)."""
        if linha.alteracao:
            return contratos
        palavras = set(chave(linha.aba).split())
        return [
            c
            for c in contratos
            if chave(c.ramo) in palavras or set(chave(c.seguradora).split()[:1]) & palavras
        ]

    def _solicitante(self, linha: Linha) -> Contato | None:
        nome = chave(linha["solicitante"])
        if not nome or not self.cliente.eh_empresa:
            return None
        achados = [c for c in self.contatos if chave(c.nome) == nome]
        if not achados:
            primeiro = nome.split()[0]
            achados = [c for c in self.contatos if chave(c.nome).split()[:1] == [primeiro]]
        if len(achados) == 1:
            return achados[0]
        situacao = "mais de um contato" if achados else "nenhum contato"
        self.avisos.append(
            f"Solicitante {linha['solicitante']}: {situacao} com este nome; fica em branco."
        )
        return None

    def _pessoa(self, linha: Linha, solicitante: Contato | None) -> Beneficiario:
        nome = chave(linha["nome"])
        if linha.dependente:
            achados = [p for p in self.pessoas if not p.eh_titular and chave(p.nome) == nome]
            if len(achados) == 1:
                return achados[0]
            if achados:
                raise ValidationError("Há mais de um dependente com este nome.")
            raise ValidationError(
                "Dependente sem cadastro: cadastre pela página do titular e importe de novo."
            )
        titulares = [p for p in self.pessoas if p.eh_titular]
        matricula = linha["cadastro"]
        if matricula:
            achados = [p for p in titulares if mesmo_codigo(p.matricula, matricula)]
            if len(achados) > 1:
                raise ValidationError(f"Há mais de um titular com a matrícula {matricula}.")
            if achados:
                pessoa = achados[0]
                if chave(pessoa.nome) != nome:
                    self.avisos.append(f"No sistema, a matrícula {matricula} é de {pessoa.nome}.")
                return pessoa
        achados = [p for p in titulares if chave(p.nome) == nome]
        if len(achados) > 1:
            raise ValidationError("Há mais de um titular com este nome: informe a matrícula.")
        if achados:
            return self._completar_matricula(achados[0], matricula)
        if linha.alteracao:
            raise ValidationError("Titular sem cadastro: na alteração, a pessoa já tem um plano.")
        return self._novo_titular(linha, matricula, solicitante)

    def _completar_matricula(self, pessoa: Beneficiario, matricula: str) -> Beneficiario:
        if matricula and pessoa.matricula and not mesmo_codigo(pessoa.matricula, matricula):
            raise ValidationError(f"No sistema, {pessoa.nome} tem a matrícula {pessoa.matricula}.")
        if matricula and not pessoa.matricula:
            pessoa.matricula = matricula
            pessoa.save(update_fields=["matricula", "atualizado_em"])
            self.avisos.append("Matrícula preenchida no cadastro da pessoa.")
        return pessoa

    def _novo_titular(self, linha: Linha, matricula: str, solicitante) -> Beneficiario:
        unidade = self._unidade_do_titular_novo(linha, solicitante)
        pessoa = Beneficiario(
            cliente=self.cliente, nome=linha["nome"], matricula=matricula, unidade=unidade
        )
        pessoa.full_clean()
        pessoa.save()
        self.pessoas.append(pessoa)
        onde = f", na unidade {unidade}" if unidade else ""
        self.avisos.append(f"Titular novo: cadastrado só com nome e matrícula{onde}.")
        return pessoa

    def _unidade_do_titular_novo(self, linha: Linha, solicitante) -> Unidade | None:
        """Com mais de uma unidade, a prévia deixa escolher, já marcando a mais provável."""
        if len(self.unidades) <= 1:
            return self.unidades[0] if self.unidades else None
        # A mesma pessoa pode vir em mais de uma aba (ex.: saúde e dental): escolhe-se uma vez.
        quem = chave(linha["cadastro"]).lstrip("0") or chave(linha["nome"])
        if quem in self.unidades_novas:
            primeira, unidade = self.unidades_novas[quem]
            if unidade is None:
                raise FaltaUnidadeError(
                    f"Titular novo: escolha a unidade na linha {primeira.numero} "
                    f"da aba {primeira.aba}."
                )
            return unidade
        self.opcoes = self.unidades
        unidade = self.unidades_escolhidas.get(linha.id) or self._unidade_provavel(
            linha, solicitante
        )
        self.unidades_novas[quem] = (linha, unidade)
        if unidade is None:
            raise FaltaUnidadeError("Titular novo: escolha a unidade em que trabalha.")
        self.escolhida = unidade.pk
        return unidade

    def _unidade_provavel(self, linha: Linha, solicitante: Contato | None) -> Unidade | None:
        """A da coluna FILIAL ou, se o solicitante atende uma unidade só, a dele."""
        filial = chave(linha["filial"])
        if filial:
            pela_filial = [
                u
                for u in self.unidades
                if filial in (chave(u), chave(u.razao_social), chave(u.nome_fantasia))
            ]
            if len(pela_filial) == 1:
                return pela_filial[0]
        do_solicitante = list(solicitante.unidades.all()) if solicitante else []
        return do_solicitante[0] if len(do_solicitante) == 1 else None

    def _avisar_se_ja_tem_plano_do_ramo(self, pessoa: Beneficiario, contrato: Contrato) -> None:
        """Inclusão de quem já tem plano do ramo pode ser troca de plano lançada na aba errada."""
        outros = pessoa.vinculos.filter(
            data_exclusao__isnull=True, contrato__ramo=contrato.ramo_id
        ).exclude(contrato=contrato)
        planos = ", ".join(v.nome_do_plano for v in outros.select_related("contrato", "plano"))
        if planos:
            self.avisos.append(
                f"A pessoa já está no plano {planos}, de {contrato.ramo}. Se é troca de plano, "
                "lance na aba ALTERAÇÕES DE PLANO."
            )

    def _plano_anterior(
        self, linha: Linha, pessoa: Beneficiario, novo: Contrato, plano: Plano | None
    ) -> Vinculo:
        """O plano atual da pessoa no mesmo ramo (ex.: PLENO, em "Pleno para Flex"), que pode
        ser outro plano do mesmo contrato (ex.: TNPQ, em "TNPQ para TNP4")."""
        atuais = [
            vinculo
            for vinculo in pessoa.vinculos.filter(data_exclusao__isnull=True).select_related(
                "contrato__ramo", "plano"
            )
            if vinculo.contrato.ramo_id == novo.ramo_id
            and (vinculo.contrato_id, vinculo.plano_id) != (novo.pk, plano.pk if plano else None)
            and (vinculo.contrato_id != novo.pk or plano is not None)
        ]
        if len(atuais) > 1:
            de = chave(linha["observacao"]).split(" PARA ")[0]
            atuais = [v for v in atuais if chave(v.nome_do_plano) == de] or atuais
        if len(atuais) == 1:
            return atuais[0]
        if atuais:
            raise ValidationError(f"A pessoa está em mais de um plano de {novo.ramo}.")
        raise ValidationError(
            f"A pessoa não está em nenhum outro plano de {novo.ramo}: lance o plano atual antes."
        )

    @staticmethod
    def _ja_registrado(pedido: Movimentacao) -> Movimentacao | None:
        """O pedido em aberto desta linha, se já está no sistema; pula o que já foi feito."""
        pessoa, contrato = pedido.beneficiario, pedido.contrato
        entrando = pessoa.movimentacoes.filter(
            contrato=contrato, tipo__in=[Movimentacao.Tipo.INCLUSAO, Movimentacao.Tipo.ALTERACAO]
        )
        em_aberto = entrando.em_aberto().first()
        if em_aberto and em_aberto.tipo == pedido.tipo:
            return em_aberto
        if em_aberto:
            raise NaoImportarError(
                f"Já há um pedido em aberto para o plano {em_aberto.nome_do_plano}."
            )
        atual = pessoa.vinculos.filter(contrato=contrato, data_exclusao__isnull=True).first()
        # Na troca dentro do mesmo contrato, a pessoa já está nele, mas em outro plano.
        trocando = pedido.eh_alteracao and atual and atual.plano_id != pedido.plano_id
        if atual and not trocando:
            raise NaoImportarError(f"Já está no plano {atual.nome_do_plano}.")
        if entrando.filter(enviada_em=pedido.enviada_em, plano=pedido.plano_id).exists():
            raise NaoImportarError("Já registrada no sistema.")
        return None
