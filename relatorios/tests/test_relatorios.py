"""Relatórios: vidas ativas, inclusões por período e apólices, na tela e em Excel."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apolices.models import Contrato
from beneficiarios.models import Beneficiario
from beneficiarios.tests.fixtures import HOJE, contrato, incluir
from clientes.models import Cliente, Contato, Unidade
from movimentacoes.models import Movimentacao
from relatorios.excel import TIPO_XLSX

pytestmark = pytest.mark.django_db

VIDAS = reverse("relatorios:vidas")
MOVIMENTACOES = reverse("relatorios:movimentacoes")
APOLICES = reverse("relatorios:apolices")


def planilha(resposta):
    assert resposta["Content-Type"] == TIPO_XLSX
    assert resposta["Cache-Control"] == "private, no-store"
    return load_workbook(BytesIO(b"".join(resposta.streaming_content))).active


def colunas(aba, *titulos) -> list[tuple]:
    """Os valores das colunas pedidas, da linha 3 em diante (1: título, 2: nomes das colunas)."""
    nomes = [celula.value for celula in aba[2]]
    posicoes = [nomes.index(titulo) for titulo in titulos]
    return [
        tuple(linha[i] for i in posicoes) for linha in aba.iter_rows(min_row=3, values_only=True)
    ]


def consultas_da_pagina(cliente_http, url: str) -> int:
    with CaptureQueriesContext(connection) as consultas:
        cliente_http.get(url)
    return len(consultas.captured_queries)


# ------------------------------------------------------------------
# Página dos relatórios
# ------------------------------------------------------------------


def test_pagina_dos_relatorios_e_menu(logado):
    resposta = logado.get(reverse("relatorios:indice"))
    pagina = resposta.content.decode()
    for url in (VIDAS, MOVIMENTACOES, APOLICES):
        assert f'href="{url}"' in pagina
    abertos = [item["texto"] for item in resposta.context["menu_principal"] if item["aberto"]]
    assert abertos == ["Relatórios"]


def test_relatorios_exigem_login(client):
    for url in (reverse("relatorios:indice"), VIDAS, MOVIMENTACOES, APOLICES):
        resposta = client.get(url)
        assert resposta.status_code == 302 and "/entrar/" in resposta.url


# ------------------------------------------------------------------
# Vidas ativas
# ------------------------------------------------------------------


def nomes(resposta) -> list[str]:
    return [vinculo.beneficiario.nome for vinculo in resposta.context["registros"]]


@pytest.fixture
def vidas(empresa, matriz, filial, saude, dental, titular, conjuge, filho):
    """Na filial: titular na saúde e no dental, cônjuge na saúde; o filho saiu em março.

    Na matriz: uma pessoa no dental e outra que só entra daqui a um mês.
    """
    incluir(titular, saude, date(2026, 1, 1))
    incluir(titular, dental, date(2026, 1, 1))
    incluir(conjuge, saude, date(2026, 1, 1))
    incluir(filho, saude, date(2026, 1, 1), exclusao=date(2026, 3, 1))
    marta = Beneficiario.objects.create(cliente=empresa, unidade=matriz, nome="Marta Matriz")
    incluir(marta, dental, date(2026, 6, 1))
    futuro = Beneficiario.objects.create(cliente=empresa, unidade=matriz, nome="Fábio Futuro")
    incluir(futuro, saude, HOJE + timedelta(days=30))


@pytest.fixture
def padaria(db, seguradora):
    """Outra empresa, com um plano de saúde e uma pessoa nele."""
    padaria = Cliente.objects.create(tipo="PJ", nome="Padaria Exemplo")
    unidade = Unidade.objects.create(
        cliente=padaria, cnpj="12ABC34501DE35", razao_social="Padaria Exemplo Ltda", matriz=True
    )
    pessoa = Beneficiario.objects.create(cliente=padaria, unidade=unidade, nome="Paula Padeira")
    incluir(pessoa, contrato(padaria, seguradora, "Saúde", "P-1", [unidade]), date(2026, 1, 1))
    return padaria


def test_vidas_ativas_hoje(logado, vidas, padaria):
    resposta = logado.get(VIDAS)
    assert resposta.context["filtro"]["data"].value() == f"{HOJE:%d/%m/%Y}"
    assert sorted(nomes(resposta)) == [
        "Ana Titular",
        "Ana Titular",
        "Bruno Cônjuge",
        "Marta Matriz",
        "Paula Padeira",
    ]
    resumo = resposta.context["resumo"]
    assert (resumo["vidas"], resumo["pessoas"]) == (5, 4)
    assert (resumo["titulares"], resumo["dependentes"]) == (3, 1)


def test_vidas_por_empresa_filial_e_contrato(logado, vidas, padaria, empresa, filial, dental):
    assert "Paula Padeira" not in nomes(logado.get(VIDAS, {"cliente": empresa.pk}))
    # O dependente fica na filial do titular, logo abaixo dele.
    da_filial = logado.get(VIDAS, {"unidade": filial.pk})
    assert nomes(da_filial) == ["Ana Titular", "Ana Titular", "Bruno Cônjuge"]
    do_dental = logado.get(VIDAS, {"contrato": dental.pk})
    assert nomes(do_dental) == ["Ana Titular", "Marta Matriz"]
    assert [linha["vidas"] for linha in do_dental.context["resumo"]["por_contrato"]] == [2]
    assert "Dental · Seguradora Exemplo D-1" in do_dental.context["descricao_do_filtro"]


def test_vidas_numa_data_passada(logado, vidas):
    resposta = logado.get(VIDAS, {"data": "15/02/2026"})
    assert sorted(nomes(resposta)) == ["Ana Titular", "Ana Titular", "Bruno Cônjuge", "Caio Filho"]


def test_vidas_em_excel(logado, vidas, titular):
    Beneficiario.objects.filter(pk=titular.pk).update(data_nascimento=date(1990, 5, 20))
    resposta = logado.get(VIDAS, {"data": f"{HOJE:%d/%m/%Y}", "formato": "excel"})
    assert f'filename="vidas-ativas-{HOJE:%Y-%m-%d}.xlsx"' in resposta["Content-Disposition"]
    aba = planilha(resposta)
    assert aba["A1"].value == f"Vidas ativas · em {HOJE:%d/%m/%Y}"
    assert aba.freeze_panes == "A3" and aba.max_row == 2 + 4
    linhas = colunas(aba, "Nome", "Titular ou dependente", "CPF", "Nascimento", "Ramo")
    ana = [linha for linha in linhas if linha[0] == "Ana Titular"]
    assert [linha[1:3] for linha in ana] == [("Titular", "529.982.247-25")] * 2
    assert ana[0][3].date() == date(1990, 5, 20)
    assert sorted(linha[4] for linha in ana) == ["Dental", "Saúde"]
    assert ("Bruno Cônjuge", "Cônjuge/companheiro(a)") in [linha[:2] for linha in linhas]


def test_vidas_consultas_nao_crescem(logado, vidas, empresa, filial, saude):
    antes = consultas_da_pagina(logado, VIDAS)
    for numero in range(5):
        pessoa = Beneficiario.objects.create(
            cliente=empresa, unidade=filial, nome=f"Pessoa {numero}"
        )
        incluir(pessoa, saude, date(2026, 2, 1))
    assert consultas_da_pagina(logado, VIDAS) == antes


# ------------------------------------------------------------------
# Inclusões por período
# ------------------------------------------------------------------


def pedido(pessoa, plano, tipo="INC", etapa="ANA", vigencia=date(2026, 10, 1), **extra):
    return Movimentacao.objects.create(
        tipo=tipo,
        beneficiario=pessoa,
        contrato=plano,
        etapa=etapa,
        recebida_em=extra.pop("recebida_em", date(2026, 9, 10)),
        enviada_em=extra.pop("enviada_em", date(2026, 9, 10)),
        vigencia=vigencia,
        **extra,
    )


def outubro(**extra) -> dict:
    return {"de": "01/10/2026", "ate": "31/10/2026", **extra}


def registros(resposta) -> list:
    return list(resposta.context["registros"])


def test_inclusoes_pela_vigencia_ou_pelo_envio(logado, titular, conjuge, filho, saude, dental):
    em_analise = pedido(titular, saude)
    concluida = pedido(conjuge, saude, etapa="CON", carteirinha="123")
    em_novembro = pedido(filho, saude, vigencia=date(2026, 11, 1), enviada_em=date(2026, 10, 20))
    pedido(titular, dental, tipo="EXC", etapa="CON")

    resposta = logado.get(MOVIMENTACOES, outubro())
    assert registros(resposta) == [em_analise, concluida]
    assert resposta.context["resumo"]["etapas"] == ["1 em análise", "1 concluída"]
    assert resposta.context["titulo"] == "Inclusões por período"
    assert "pela vigência" in resposta.context["descricao_do_filtro"]

    pelo_envio = logado.get(MOVIMENTACOES, outubro(contar_por="envio"))
    assert registros(pelo_envio) == [em_novembro]
    assert "pelo envio à operadora" in pelo_envio.context["descricao_do_filtro"]


def test_tipo_situacao_e_contrato(
    logado, titular, conjuge, filho, saude, dental, empresa, seguradora, matriz, filial
):
    inclusao = pedido(titular, saude)
    rejeitada = pedido(filho, saude, etapa="REJ", motivo_rejeicao="RAS")
    exclusao = pedido(conjuge, dental, tipo="EXC", etapa="CON")
    outro_saude = contrato(empresa, seguradora, "Saúde", "S-2", [matriz, filial])
    alteracao = pedido(titular, outro_saude, tipo="ALT", contrato_anterior=saude)

    exclusoes = logado.get(MOVIMENTACOES, outubro(tipo="EXC"))
    assert registros(exclusoes) == [exclusao]
    assert exclusoes.context["titulo"] == "Exclusões por período"
    todos = logado.get(MOVIMENTACOES, outubro(tipo="todos"))
    assert set(registros(todos)) == {inclusao, rejeitada, exclusao, alteracao}
    assert todos.context["titulo"] == "Movimentações por período"
    abertas = logado.get(MOVIMENTACOES, outubro(tipo="todos", situacao="aberto"))
    assert set(registros(abertas)) == {inclusao, alteracao}
    assert registros(logado.get(MOVIMENTACOES, outubro(situacao="REJ"))) == [rejeitada]
    # A alteração aparece no plano de onde a pessoa sai e no plano para onde vai.
    do_plano = logado.get(MOVIMENTACOES, outubro(tipo="todos", contrato=saude.pk))
    assert set(registros(do_plano)) == {inclusao, rejeitada, alteracao}


def test_periodo_invalido_e_padrao(logado):
    resposta = logado.get(MOVIMENTACOES, {"de": "31/10/2026", "ate": "01/10/2026"})
    assert "Escolha uma data igual ou depois da primeira." in resposta.content.decode()
    assert "resumo" not in resposta.context
    # Sem filtro, abre no mês atual.
    padrao = logado.get(MOVIMENTACOES)
    assert padrao.context["filtro"]["de"].value() == f"{HOJE.replace(day=1):%d/%m/%Y}"
    assert padrao.context["resumo"]["total"] == 0


def test_inclusoes_em_excel(logado, titular, saude, empresa, seguradora, matriz, filial):
    formula = '=HYPERLINK("https://exemplo.invalid","abrir")'
    pedido(titular, saude, observacao=formula)
    pleno = contrato(empresa, seguradora, "Saúde", "S-2", [matriz, filial])
    Contrato.objects.filter(pk=pleno.pk).update(produto="PLENO")
    pedido(titular, pleno, tipo="ALT", contrato_anterior=saude)

    aba = planilha(logado.get(MOVIMENTACOES, outubro(tipo="todos", formato="excel")))
    assert aba["A1"].value == "Movimentações por período · 01/10/2026 a 31/10/2026, pela vigência"
    assert sorted(colunas(aba, "Tipo", "Plano")) == [
        ("Alteração de plano", "S-1 para PLENO"),
        ("Inclusão", "S-1"),
    ]
    assert {valor.date() for (valor,) in colunas(aba, "Vigência")} == {date(2026, 10, 1)}
    # Texto que parece fórmula fica como texto.
    posicao = [celula.value for celula in aba[2]].index("Observação") + 1
    celulas = [aba.cell(row=linha, column=posicao) for linha in range(3, aba.max_row + 1)]
    [observacao] = [celula for celula in celulas if celula.value == formula]
    assert observacao.data_type == "s"


# ------------------------------------------------------------------
# Apólices
# ------------------------------------------------------------------


def numeros(resposta) -> list[str]:
    return [contrato.numero for contrato, _ in resposta.context["registros"]]


def test_apolices_pela_renovacao(logado, empresa, seguradora, matriz, saude, dental, vida, titular):
    Contrato.objects.filter(pk=saude.pk).update(
        data_renovacao=date(2026, 11, 1), valor_mensal=Decimal("1000.50")
    )
    Contrato.objects.filter(pk=dental.pk).update(data_renovacao=date(2026, 10, 15))
    cancelado = contrato(empresa, seguradora, "Saúde", "S-9", [matriz])
    Contrato.objects.filter(pk=cancelado.pk).update(
        situacao="CAN", data_renovacao=date(2026, 10, 20)
    )
    incluir(titular, saude, date(2026, 1, 1))

    # Vigentes, pela data de renovação; os sem data ficam no fim.
    resposta = logado.get(APOLICES)
    assert numeros(resposta) == ["D-1", "S-1", "V-1"]
    assert resposta.context["resumo"] == {"total": 3, "valor": Decimal("1000.50"), "vidas": 1}
    periodo = {"renovacao_de": "01/10/2026", "renovacao_ate": "31/10/2026"}
    assert numeros(logado.get(APOLICES, periodo)) == ["D-1"]
    assert numeros(logado.get(APOLICES, {**periodo, "situacao": "todas"})) == ["D-1", "S-9"]
    assert numeros(logado.get(APOLICES, {"ramo": saude.ramo_id})) == ["S-1"]


def test_dados_do_cliente(logado, empresa, matriz, filial, seguradora, saude, contato_rh, pessoa):
    Contato.objects.filter(pk=contato_rh.pk).update(
        telefone="(11) 3333-4444", email="rh@exemplo.com.br"
    )
    Unidade.objects.filter(pk=matriz.pk).update(
        cidade="Curitiba", uf="PR", telefone="(41) 3000-0000"
    )
    Cliente.objects.filter(pk=pessoa.pk).update(
        telefone="(11) 98888-7777", email="joao@exemplo.com.br", cidade="Santos", uf="SP"
    )
    Contrato.objects.create(
        cliente=pessoa,
        seguradora=seguradora,
        ramo=saude.ramo,
        numero="IND-1",
        modalidade="IF",
        inicio_vigencia=date(2026, 1, 1),
        dia_vencimento=5,
    )

    def dados() -> dict:
        return {c.numero: d for c, d in logado.get(APOLICES).context["registros"]}

    assert dados()["S-1"] == {
        "documento": "11.222.333/0001-81",
        "contato": "Carla RH",
        "telefone": "(11) 3333-4444",
        "email": "rh@exemplo.com.br",
        "cidade": "Curitiba/PR",
    }
    assert dados()["IND-1"] == {
        "documento": "529.982.247-25",
        "contato": "",
        "telefone": "(11) 98888-7777",
        "email": "joao@exemplo.com.br",
        "cidade": "Santos/SP",
    }
    # Sem contato principal, valem o telefone e o e-mail da matriz.
    Contato.objects.filter(pk=contato_rh.pk).update(principal=False)
    assert dados()["S-1"]["telefone"] == "(41) 3000-0000"


def test_apolices_em_excel(logado, empresa, saude, titular, contato_rh):
    Contrato.objects.filter(pk=saude.pk).update(
        data_renovacao=date(2026, 11, 1), valor_mensal=Decimal("18450.90")
    )
    incluir(titular, saude, date(2026, 1, 1))
    aba = planilha(logado.get(APOLICES, {"formato": "excel"}))
    assert aba["A1"].value == "Apólices · vigentes"
    [linha] = colunas(aba, "Cliente", "CPF/CNPJ", "Contato", "Renovação", "Valor mensal", "Vidas")
    assert linha[:3] == ("Supermercados Exemplo", "11.222.333/0001-81", "Carla RH")
    assert linha[3].date() == date(2026, 11, 1)
    assert linha[4] == pytest.approx(18450.90) and linha[5] == 1
    posicao = [celula.value for celula in aba[2]].index("Valor mensal") + 1
    assert aba.cell(row=3, column=posicao).number_format == "#,##0.00"


def test_apolices_consultas_nao_crescem(logado, empresa, seguradora, matriz):
    contrato(empresa, seguradora, "Saúde", "A-1", [matriz])
    antes = consultas_da_pagina(logado, APOLICES)
    for numero in range(2, 7):
        outra = Cliente.objects.create(tipo="PJ", nome=f"Empresa {numero}")
        contrato(outra, seguradora, "Dental", f"A-{numero}", [])
    assert consultas_da_pagina(logado, APOLICES) == antes
