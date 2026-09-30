"""Planilhas do Excel dos relatórios, no estilo das que o sistema já gera.

Linha 1: o título (com os filtros usados); linha 2: os nomes das colunas; os dados a partir
da linha 3, com filtro nas colunas e o cabeçalho congelado.
"""

from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO

from django.http import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATOS = {"data": "DD/MM/YYYY", "dinheiro": "#,##0.00"}


@dataclass(frozen=True)
class Coluna:
    titulo: str
    largura: int
    # "data" ou "dinheiro"; em branco, o valor vai como está.
    formato: str = ""


def gerar(titulo: str, aba: str, colunas: list[Coluna], linhas: Iterable[list]) -> bytes:
    livro = Workbook()
    folha = livro.active
    folha.title = aba[:31]
    folha.append([titulo])
    folha.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colunas))
    folha["A1"].font = Font(bold=True, size=12)
    folha.append([coluna.titulo for coluna in colunas])
    for celula in folha[2]:
        celula.font = Font(bold=True)
    for linha in linhas:
        folha.append(linha)
    for numero, coluna in enumerate(colunas, start=1):
        folha.column_dimensions[get_column_letter(numero)].width = coluna.largura
        formato = FORMATOS.get(coluna.formato)
        for (celula,) in folha.iter_rows(min_row=3, min_col=numero, max_col=numero):
            if formato:
                celula.number_format = formato
            elif isinstance(celula.value, str) and celula.value.startswith("="):
                # Texto que começa com "=" (ex.: numa observação) fica como texto, nunca fórmula.
                celula.data_type = "s"
    folha.freeze_panes = "A3"
    folha.auto_filter.ref = f"A2:{get_column_letter(len(colunas))}{max(folha.max_row, 2)}"
    arquivo = BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def baixar(conteudo: bytes, nome_do_arquivo: str) -> FileResponse:
    resposta = FileResponse(
        BytesIO(conteudo), as_attachment=True, filename=nome_do_arquivo, content_type=TIPO_XLSX
    )
    # Os relatórios têm dados pessoais: o navegador não guarda cópia.
    resposta["Cache-Control"] = "private, no-store"
    return resposta
